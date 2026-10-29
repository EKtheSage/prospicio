"""Writes a test workbook for the XLL: RAA with every function in place
and R ChainLadder's totals beside them.

    python3 make_workbook.py prospicio-xll-test.xlsx

Each table is saved as an Excel 365 dynamic-array formula that spills from
its first cell, as if typed there. openpyxl writes only legacy (Ctrl+Shift+
Enter) array formulas, so the saved file is patched: each array formula's
cell gets `cm="1"`, pointing at a dynamic-array entry in `xl/metadata.xml`,
which is how Excel itself saves a spilling formula.
"""

import csv
import datetime
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.styles import Font, PatternFill
out = sys.argv[1]
RAA = [
 [5012,8269,10907,11805,13539,16181,18009,18608,18662,18834],
 [106,4285,5396,10666,13782,15599,15496,16169,16704],
 [3410,8992,13873,16141,18735,22214,22863,23466],
 [5655,11555,15766,21266,23425,26083,27067],
 [1092,9565,15836,22169,25955,26180],
 [1513,6445,11702,12935,15852],
 [557,4020,10946,12314],
 [1351,6947,13112],
 [3133,5395],
 [2063]]
wb = Workbook(); ws = wb.active; ws.title = "RAA"
bold = Font(bold=True); note = Font(italic=True, color="666666")
inp = PatternFill("solid", fgColor="FFF2CC")
ws["A1"] = "prospicio XLL test: RAA (Mack 1993), paid, cumulative"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "Load prospicio.xll first (File > Options > Add-ins > Excel Add-ins > Browse). If cells show #NAME?, press Ctrl+Alt+F9."; ws["A2"].font = note
ws["M1"] = "Add-in:"; ws["N1"] = ArrayFormula("N1", "=PRO.VERSION()")
ws["A3"] = "AY"; ws["A3"].font = bold
for c in range(10):
    cell = ws.cell(row=3, column=2+c, value=12*(c+1)); cell.font = bold
for r, row in enumerate(RAA):
    ws.cell(row=4+r, column=1, value=1981+r).font = bold
    for c, x in enumerate(row):
        ws.cell(row=4+r, column=2+c, value=x).fill = inp
T = "B4:K13"
ws["A14"] = "LDF"; ws["A14"].font = bold
ws["B14"] = ArrayFormula("B14:K14", f"=PRO.CL.LDF({T})")
ws["A15"] = "CDF"; ws["A15"].font = bold
ws["B15"] = ArrayFormula("B15:K15", f"=PRO.CL.CDF({T})")
ws["A16"] = "LDF, simple avg, tail 1.05"; ws["A16"].font = note
ws["B16"] = ArrayFormula("B16:K16", f'=PRO.CL.LDF({T},"simple",1.05)')
ws["A17"] = "LDF, latest 5, ex high/low"; ws["A17"].font = note
ws["B17"] = ArrayFormula("B17:K17", f'=PRO.CL.LDF({T},{{"volume",5,1,1}})')

ws["M3"] = ArrayFormula("M3:Q14", f"=PRO.CL({T},,,A4:A13)")
ws["M16"] = "PRO.CL.ULTIMATE and PRO.CL.RESERVE as plain columns:"; ws["M16"].font = note
ws["M17"] = ArrayFormula("M17:M26", f"=PRO.CL.ULTIMATE({T})")
ws["N17"] = ArrayFormula("N17:N26", f"=PRO.CL.RESERVE({T})")

ws["A18"] = "Mack (volume-weighted, no tail)"; ws["A18"].font = bold
ws["A19"] = ArrayFormula("A19:H30", f"=PRO.MACK({T},,,A4:A13)")
ws["A32"] = "Development factors, exponential tail"; ws["A32"].font = bold
ws["A33"] = ArrayFormula("A33:E43", f'=PRO.CL.FACTORS({T},,"exponential")')

ws["J32"] = "Check against R ChainLadder 0.2.21"; ws["J32"].font = bold
checks = [
 ("Total CL reserve", "=Q14", 52135.2282612101),
 ("Total Mack SE", "=G30", 26880.7403298914),
 ("Total process SE", "=E30", 24917.2657465242),
 ("Total parameter SE", "=F30", 10084.8434990404),
]
for h, col in zip(["Quantity","prospicio","R","Match"], "JKLM"):
    ws[f"{col}33"] = h; ws[f"{col}33"].font = bold
for i, (name, f, r) in enumerate(checks):
    row = 34+i
    ws[f"J{row}"] = name; ws[f"K{row}"] = f; ws[f"L{row}"] = r
    ws[f"M{row}"] = f"=ABS(K{row}-L{row})<1E-6"
ws["J39"] = "Try an error: average \"median\""; ws["J39"].font = note
ws["J40"] = ArrayFormula("J40", f'=PRO.CL.LDF({T},"median")')
for col, w in {"A":12,"J":22,"M":12,"N":12}.items():
    ws.column_dimensions[col].width = w
for col in "BCDEFGHIKLOPQ":
    ws.column_dimensions[col].width = 12
for row in ws.iter_rows(min_row=4, max_row=43):
    for c in row:
        if c.column > 1 and c.number_format == "General":
            c.number_format = "#,##0.0000" if c.row in (14,15,16) or (33 <= c.row <= 43 and c.column <= 5) else "#,##0"
# Layer pricing: parameters in cells, a column of attachments against a
# row of limits, so one formula fills the table.
ly = wb.create_sheet("Layers")
ly["A1"] = "prospicio XLL test: layer pricing"; ly["A1"].font = Font(bold=True, size=13)
ly["A2"] = "Change the yellow cells; every table recalculates."; ly["A2"].font = note
for row, (label, value) in enumerate([
    ("Pareto t", 500), ("Pareto alpha", 2),
    ("Lognormal mean", 50000), ("Lognormal CV", 3), ("Basic limit", 100000)], start=4):
    ly[f"A{row}"] = label; ly[f"A{row}"].font = bold
    ly[f"B{row}"] = value; ly[f"B{row}"].fill = inp
PAR = '"pareto",$B$4,$B$5'
LOG = '"lognormal_cv",$B$6,$B$7'
ly["A10"] = "Pareto: expected loss per ground-up loss, PRO.LAYER"; ly["A10"].font = bold
ly["A11"] = "Attachment \\ limit"; ly["A11"].font = note
for c, lim in enumerate([1000, 4000, 10000, "unlimited"]):
    cell = ly.cell(row=11, column=2+c, value=lim); cell.fill = inp; cell.font = bold
for r, a in enumerate([1000, 5000, 10000]):
    cell = ly.cell(row=12+r, column=1, value=a); cell.fill = inp; cell.font = bold
ly["B12"] = ArrayFormula("B12:E14", f"=PRO.LAYER(B11:E11,A12:A14,{PAR})")
ly["A16"] = "Layer CV given a loss above the attachment, PRO.LAYER.CV"; ly["A16"].font = note
ly["B17"] = ArrayFormula("B17:E19", f"=PRO.LAYER.CV(B11:E11,A12:A14,{PAR})")

ly["G10"] = "Lognormal: ILF and LEV by limit"; ly["G10"].font = bold
for col, h in zip("GHIJ", ["Limit", "ILF", "LEV", "LEV, Pareto"]):
    ly[f"{col}11"] = h; ly[f"{col}11"].font = bold
for r, lim in enumerate([100000, 250000, 500000, 1000000, 2000000, 5000000, "unlimited"]):
    ly.cell(row=12+r, column=7, value=lim).fill = inp
ly["H12"] = ArrayFormula("H12:H18", f"=PRO.ILF(G12:G18,$B$8,{LOG})")
ly["I12"] = ArrayFormula("I12:I18", f"=PRO.LEV(G12:G18,{LOG})")
ly["J12"] = ArrayFormula("J12:J18", f"=PRO.LEV(G12:G18,{PAR})")

ly["A22"] = "Check against the R package Pareto (vignette, t = 500, alpha = 2)"; ly["A22"].font = bold
for h, col in zip(["Quantity", "prospicio", "Expected", "Match"], "ABCD"):
    ly[f"{col}23"] = h; ly[f"{col}23"].font = bold
for i, (name, f, x) in enumerate([
        ("4000 xs 1000", "=C12", 200),
        ("unlimited xs 1000, t^a d^(1-a)/(a-1)", "=E12", 250)]):
    row = 24 + i
    ly[f"A{row}"] = name; ly[f"B{row}"] = f; ly[f"C{row}"] = x
    ly[f"D{row}"] = f"=ABS(B{row}-C{row})<1E-9"
ly["A27"] = "Valid only with the inputs above unchanged."; ly["A27"].font = note

ly["A30"] = "Every severity family: mean, 99.5% quantile and LEV at 1,000,000"; ly["A30"].font = bold
for h, col in zip(["Family", "p1", "p2", "p3", "Mean", "Q 99.5%", "LEV 1m", "LER at 10,000"], "ABCDEFGH"):
    ly[f"{col}31"] = h; ly[f"{col}31"].font = bold
for i, (fam, ps) in enumerate([
        ("pareto", [500, 2]), ("lognormal_cv", [50000, 3]), ("gamma_cv", [50000, 1]),
        ("weibull", [0.7, 30000]), ("loglogistic", [2.5, 30000]), ("inverse_gamma", [3, 100000]),
        ("inverse_gaussian", [50000, 20000]), ("burr", [3, 1.5, 60000]), ("gpd", [0.3, 30000])]):
    r = 32 + i
    ly[f"A{r}"] = fam
    for c, x in zip("BCD", ps):
        ly[f"{c}{r}"] = x; ly[f"{c}{r}"].fill = inp
    a = f'A{r},B{r},C{r},D{r}'
    ly[f"E{r}"] = ArrayFormula(f"E{r}", f"=INDEX(PRO.SEV.STATS({a}),2,2)")
    ly[f"F{r}"] = ArrayFormula(f"F{r}", f"=PRO.SEV.QUANTILE(0.995,{a})")
    ly[f"G{r}"] = ArrayFormula(f"G{r}", f"=PRO.LEV(1000000,{a})")
    ly[f"H{r}"] = ArrayFormula(f"H{r}", f"=PRO.LER(10000,{a})")
    for c in "EFG":
        ly[f"{c}{r}"].number_format = "#,##0"
    ly[f"H{r}"].number_format = "0.0000"

ly["A43"] = "Pareto helpers (R package Pareto)"; ly["A43"].font = bold
ly["A44"] = "Alpha from two layers: 1m xs 1m expects 1, 2m xs 3m expects"; ly["F44"] = 0.6; ly["F44"].fill = inp
ly["G44"] = ArrayFormula("G44", "=PRO.PARETO.ALPHA(1000000,1000000,1,2000000,3000000,F44)")
ly["A45"] = "Alpha from 4 losses above 1m and 1 above 2m"
ly["G45"] = ArrayFormula("G45", "=PRO.PARETO.ALPHA.FREQS(1000000,4,2000000,1)")
ly["A46"] = "1m xs 1m at alpha G44, extrapolated to each layer below (per unit)"
for c, lim in enumerate([1000000, 2000000, 5000000]):
    ly.cell(row=47, column=2+c, value=lim).fill = inp
for r, att in enumerate([2000000, 3000000]):
    ly.cell(row=48+r, column=1, value=att).fill = inp
ly["B48"] = ArrayFormula("B48:D49", "=PRO.PARETO.EXTRAPOLATE(1000000,1000000,B47:D47,A48:A49,G44)")

ly["A52"] = "Piecewise Pareto: thresholds and alphas as rows, truncated at 50m (last piece)"; ly["A52"].font = bold
for c, x in zip("BCD", [1000000, 2000000, 5000000]):
    ly[f"{c}53"] = x; ly[f"{c}53"].fill = inp
for c, x in zip("BCD", [1.2, 1.8, 2.5]):
    ly[f"{c}54"] = x; ly[f"{c}54"].fill = inp
ly["A53"] = "Thresholds"; ly["A54"] = "Alphas"; ly["A55"] = "Truncation"
ly["B55"] = 50000000; ly["B55"].fill = inp; ly["C55"] = "lp"; ly["C55"].fill = inp
ly["A56"] = "Layer"; ly["B56"] = "Expected loss"; ly["C56"] = "SD"
for h in ("A56", "B56", "C56"):
    ly[h].font = bold
for i, (lim, att) in enumerate([(1000000, 1000000), (3000000, 2000000), (5000000, 5000000)]):
    r = 57 + i
    ly[f"A{r}"] = f"{lim // 1000000}m xs {att // 1000000}m"
    a = f'{lim},{att},"piecewise_pareto",$B$53:$D$53,$B$54:$D$54,$B$55:$C$55'
    ly[f"B{r}"] = ArrayFormula(f"B{r}", f"=PRO.LAYER({a})")
    ly[f"C{r}"] = ArrayFormula(f"C{r}", f"=SQRT(PRO.LAYER.VAR({a}))")
    for c in "BC":
        ly[f"{c}{r}"].number_format = "#,##0"
ly.column_dimensions["A"].width = 34
for col in "BCDEGHIJ":
    ly.column_dimensions[col].width = 14
for row in ly.iter_rows(min_row=12, max_row=19, min_col=2, max_col=10):
    for c in row:
        c.number_format = "0.0000" if (c.row >= 17 and c.column <= 5) or c.column == 8 else "#,##0.00"
for row in ly.iter_rows(min_row=11, max_row=18, min_col=7, max_col=7):
    for c in row:
        c.number_format = "#,##0"

# ODP bootstrap on the RAA sheet's triangle.
od = wb.create_sheet("ODP")
od["A1"] = "prospicio XLL test: ODP bootstrap on RAA"; od["A1"].font = Font(bold=True, size=13)
od["A2"] = "The triangle is on the RAA sheet. Change the yellow cells; the same seed gives the same table."
od["A2"].font = note
for row, (label, value) in enumerate([("Simulations", 20000), ("Seed", 1)], start=4):
    od[f"A{row}"] = label; od[f"A{row}"].font = bold
    od[f"B{row}"] = value; od[f"B{row}"].fill = inp
od["D4"] = "Percentiles"; od["D4"].font = bold
for c, p in enumerate([0.75, 0.9, 0.95, 0.995]):
    cell = od.cell(row=4, column=5+c, value=p); cell.fill = inp
TRI, ORI = "RAA!$B$4:$K$13", "RAA!$A$4:$A$13"
od["A7"] = "Gamma process error (R's process.distr = 'gamma')"; od["A7"].font = bold
od["A8"] = ArrayFormula("A8:I19", f"=PRO.ODP({TRI},$B$4,$B$5,$E$4:$H$4,,,{ORI})")
od["A21"] = "Parameter error only (process \"none\")"; od["A21"].font = bold
od["A22"] = ArrayFormula("A22:I33", f'=PRO.ODP({TRI},$B$4,$B$5,$E$4:$H$4,,"none",{ORI})')

od["K7"] = "Check: total against R ChainLadder 0.2.21 BootChainLadder (100,000 sims)"; od["K7"].font = bold
for h, col in zip(["Quantity", "prospicio", "R", "Tolerance", "Within"], "KLMNO"):
    od[f"{col}8"] = h; od[f"{col}8"].font = bold
checks = [
    ("Gamma: mean", "=C19", 53842.2718407722, 589.413),
    ("Gamma: SD", "=D19", 18965.3468418162, 483.93),
    ("Gamma: P95", "=H19", 87856.1780690262, 1854.44),
    ("None: mean", "=C33", 53928.2509498396, 506.674),
    ("None: SD", "=D33", 17453.9114703321, 513.428),
    ("None: P95", "=H33", 85236.8809859552, 1701.59),
]
for i, (name, f, r, tol) in enumerate(checks):
    row = 9 + i
    od[f"K{row}"] = name; od[f"L{row}"] = f; od[f"M{row}"] = r; od[f"N{row}"] = tol
    od[f"O{row}"] = f"=ABS(L{row}-M{row})<N{row}"
od["K16"] = "Tolerance: four standard errors at 20,000 simulations; use 20000 or more."
od["K16"].font = note
od.column_dimensions["A"].width = 12
od.column_dimensions["K"].width = 16
for col in "BCDEFGHILMNO":
    od.column_dimensions[col].width = 12
for row in od.iter_rows(min_row=9, max_row=33, min_col=2, max_col=14):
    for c in row:
        if not (c.column >= 11 and c.row > 14):
            c.number_format = "0.000" if c.column == 5 and c.row != 4 else "#,##0"
for c in od[4][4:8]:
    c.number_format = "0.0%"

# The draws behind the ODP sheet's gamma table: same simulations and seed.
dr = wb.create_sheet("ODP draws")
dr["A1"] = "prospicio XLL test: the ODP bootstrap's draws (PRO.ODP.SIMS)"; dr["A1"].font = Font(bold=True, size=13)
dr["A2"] = "Simulations and seed come from the ODP sheet; the average of each column is that sheet's Mean."
dr["A2"].font = note
dr["N3"] = "Check"; dr["N3"].font = bold
dr["N4"] = "Mean total, here"; dr["O4"] = "=AVERAGE(K5:K100004)"
dr["N5"] = "Mean total, ODP sheet"; dr["O5"] = "=ODP!C19"
dr["N6"] = "Match"; dr["O6"] = "=ABS(O4-O5)<1E-6*O5"
dr["N8"] = "Total reserve histogram"; dr["N8"].font = bold
dr["N9"] = "Up to"; dr["O9"] = "Count"
for h in ("N3", "N9", "O9"):
    dr[h].font = bold
for i, b in enumerate(range(20000, 160001, 20000)):
    dr.cell(row=10+i, column=14, value=b).fill = inp
dr["N18"] = "more"
# FREQUENCY returns one count per bin and one for everything above.
dr["O10"] = ArrayFormula("O10:O18", "=FREQUENCY(K5:K100004,N10:N17)")
dr["A4"] = ArrayFormula("A4:K20004", f"=PRO.ODP.SIMS({TRI},ODP!$B$4,ODP!$B$5,,,{ORI})")
dr.column_dimensions["N"].width = 22
for col in "ABCDEFGHIJKO":
    dr.column_dimensions[col].width = 11
for row in dr.iter_rows(min_row=5, max_row=200, min_col=1, max_col=11):
    for c in row:
        c.number_format = "#,##0"
dr["O4"].number_format = dr["O5"].number_format = "#,##0.00"

# Mack's bootstrap on the RAA sheet's triangle, beside Mack's closed form.
mk = wb.create_sheet("Mack bootstrap")
mk["A1"] = "prospicio XLL test: Mack's bootstrap on RAA (England, Verrall and Wuthrich)"
mk["A1"].font = Font(bold=True, size=13)
mk["A2"] = "Centred residuals: the simulated SD is Mack's standard error (the RAA sheet's PRO.MACK)."
mk["A2"].font = note
for row, (label, value) in enumerate([("Simulations", 20000), ("Seed", 1)], start=4):
    mk[f"A{row}"] = label; mk[f"A{row}"].font = bold
    mk[f"B{row}"] = value; mk[f"B{row}"].fill = inp
mk["A7"] = "Gamma process error"; mk["A7"].font = bold
mk["A8"] = ArrayFormula("A8:I19", f"=PRO.MACK.BOOT({TRI},$B$4,$B$5,,,,,{ORI})")
mk["K7"] = "Check: total against Mack's closed form"; mk["K7"].font = bold
for h, col in zip(["Quantity", "Bootstrap", "Mack", "Tolerance", "Within"], "KLMNO"):
    mk[f"{col}8"] = h; mk[f"{col}8"].font = bold
# Mack's SD with the parameter risk scaled by the residual pool's variance
# 0.981 on RAA (validation/tests/reserving_mack_bootstrap.rs).
mk["K9"] = "Mean"; mk["L9"] = "=C19"; mk["M9"] = "=RAA!D30"; mk["N9"] = "=5*D19/SQRT($B$4)"
mk["K10"] = "SD"; mk["L10"] = "=D19"; mk["M10"] = "=SQRT(RAA!E30^2+0.981*RAA!F30^2)"
mk["N10"] = "=5*D19*SQRT(5/(4*$B$4))"
for r in (9, 10):
    mk[f"O{r}"] = f"=ABS(L{r}-M{r})<N{r}"
mk.column_dimensions["K"].width = 14
for col in "ABCDEFGHILMNO":
    mk.column_dimensions[col].width = 12
for row in mk.iter_rows(min_row=9, max_row=19, min_col=2, max_col=14):
    for c in row:
        c.number_format = "0.000" if c.column == 5 else "#,##0"

# The expected-loss methods on GenIns with R ChainLadder's ClarkCapeCod
# premium, against chainladder-python 0.10.1's totals.
GENINS = [
    [357848, 1124788, 1735330, 2218270, 2745596, 3319994, 3466336, 3606286, 3833515, 3901463],
    [352118, 1236139, 2170033, 3353322, 3799067, 4120063, 4647867, 4914039, 5339085],
    [290507, 1292306, 2218525, 3235179, 3985995, 4132918, 4628910, 4909315],
    [310608, 1418858, 2195047, 3757447, 4029929, 4381982, 4588268],
    [443160, 1136350, 2128333, 2897821, 3402672, 3873311],
    [396132, 1333217, 2180715, 2985752, 3691712],
    [440832, 1288463, 2419861, 3483130],
    [359480, 1421128, 2864498],
    [376686, 1363294],
    [344014],
]
el = wb.create_sheet("Expected loss")
el["A1"] = "prospicio XLL test: expected-loss methods on GenIns (premium as in R's ?ClarkCapeCod)"
el["A1"].font = Font(bold=True, size=13)
el["A3"] = "AY"; el["L3"] = "Premium"; el["A3"].font = el["L3"].font = bold
for c in range(10):
    el.cell(row=3, column=2+c, value=12*(c+1)).font = bold
for r, row in enumerate(GENINS):
    el.cell(row=4+r, column=1, value=2001+r).font = bold
    for c, x in enumerate(row):
        el.cell(row=4+r, column=2+c, value=x).fill = inp
    el.cell(row=4+r, column=12, value=10_000_000 + 400_000*r).fill = inp
el["N3"] = "A priori LR"; el["N3"].font = bold; el["O3"] = 0.75; el["O3"].fill = inp
G, P, O = "$B$4:$K$13", "$L$4:$L$13", "$A$4:$A$13"
el["A16"] = "Bornhuetter-Ferguson"; el["A16"].font = bold
el["A17"] = ArrayFormula("A17:G28", f"=PRO.BF({G},{P},$O$3,,,{O})")
el["I16"] = "Benktander, 2 iterations"; el["I16"].font = bold
el["I17"] = ArrayFormula("I17:O28", f"=PRO.BENKTANDER({G},{P},$O$3,2,,,{O})")
el["A31"] = "Cape Cod, trend 5%, decay 0.75"; el["A31"].font = bold
el["A32"] = ArrayFormula("A32:G43", f"=PRO.CAPECOD({G},{P},0.05,0.75,,,{O})")
el["I31"] = "Expected loss"; el["I31"].font = bold
el["I32"] = ArrayFormula("I32:O43", f"=PRO.EL({G},{P},$O$3,,,{O})")
el["A46"] = "Check: total reserve against chainladder-python 0.10.1 (with a priori 0.75)"; el["A46"].font = bold
for h, col in zip(["Method", "prospicio", "chainladder-python", "Match"], "ABCD"):
    el[f"{col}47"] = h; el[f"{col}47"].font = bold
for i, (name, f, x) in enumerate([
        ("Bornhuetter-Ferguson", "=G28", 32357669.898322247),
        ("Benktander, 2", "=O28", 27674175.15049441),
        ("Cape Cod 5%, 0.75", "=G43", 22446795.86898648),
        ("Expected loss", "=O43", 54141910.0)]):
    row = 48 + i
    el[f"A{row}"] = name; el[f"B{row}"] = f; el[f"C{row}"] = x
    el[f"D{row}"] = f"=ABS(B{row}-C{row})<1E-6*C{row}"
el["A54"] = "Clark's Cape Cod, log-logistic, max age 240 (R's ClarkCapeCod)"; el["A54"].font = bold
el["A55"] = ArrayFormula("A55:G66", f"=PRO.CLARK({G},{P},,240,{O})")
el["I54"] = "Its parameters"; el["I54"].font = bold
el["I55"] = ArrayFormula("I55:J60", f"=PRO.CLARK.PARAMS({G},{P},,240)")
el["A69"] = "Clark's LDF method on GenIns, Weibull"; el["A69"].font = bold
el["A70"] = ArrayFormula("A70:G81", f'=PRO.CLARK({G},,"weibull",,{O})')
# A selected pattern at its ages and an a priori per origin (decision 40).
el["M3"] = "A priori by AY"; el["M3"].font = bold
for r in range(10):
    el.cell(row=4+r, column=13, value=round(0.70 + 0.01*r, 2)).fill = inp
el["A84"] = "GenIns's own factors to ultimate (PRO.CL.CDF) at their ages, as a selected pattern"; el["A84"].font = bold
el["A85"] = "Pattern age"; el["A86"] = "CDF"; el["A85"].font = el["A86"].font = bold
for c in range(10):
    el.cell(row=85, column=2+c, value=12*(c+1)).fill = inp
el["B86"] = ArrayFormula("B86:K86", f"=PRO.CL.CDF({G})")
el["A88"] = "Bornhuetter-Ferguson on the selected pattern"; el["A88"].font = bold
el["A89"] = ArrayFormula("A89:G100", f'=PRO.BF({G},{P},$O$3,,,{O},$B$86:$K$86,"cdf",,$B$85:$K$85)')
el["I88"] = "Bornhuetter-Ferguson, a priori by AY (column M)"; el["I88"].font = bold
el["I89"] = ArrayFormula("I89:O100", f"=PRO.BF({G},{P},$M$4:$M$13,,,{O})")
el["A102"] = "Check: the selected pattern and the a priori by AY"; el["A102"].font = bold
for h, col in zip(["Check", "prospicio", "Expected", "Match"], "ABCD"):
    el[f"{col}103"] = h; el[f"{col}103"].font = bold
for i, (name, f, x) in enumerate([
        ("Own pattern selected = BF", "=G100", "=G28"),
        ("By AY = latest + LR x premium x (1 - 1/CDF)", "=N100",
         "=SUMPRODUCT(K90:K99+M90:M99*J90:J99*(1-1/L90:L99))")]):
    row = 104 + i
    el[f"A{row}"] = name; el[f"B{row}"] = f; el[f"C{row}"] = x
    el[f"D{row}"] = f"=ABS(B{row}-C{row})<1E-6*ABS(C{row})"
el.column_dimensions["A"].width = 22
for col in "BCDEFGHIJKLMNO":
    el.column_dimensions[col].width = 12
for row in el.iter_rows(min_row=17, max_row=105, min_col=2, max_col=15):
    for c in row:
        if c.row not in (85, 86):
            c.number_format = "#,##0"
for c in range(10):
    el.cell(row=86, column=2+c).number_format = "0.0000"

# The one-year view on the RAA sheet's triangle.
oy = wb.create_sheet("One-year")
oy["A1"] = "prospicio XLL test: the one-year view on RAA"; oy["A1"].font = Font(bold=True, size=13)
oy["A2"] = "CDR = opening ultimate less the re-estimated one: a negative CDR is adverse."
oy["A2"].font = note
oy["A4"] = "Merz-Wuthrich (closed form)"; oy["A4"].font = bold
oy["A5"] = ArrayFormula("A5:D16", f"=PRO.MACK.CDR({TRI},,,{ORI})")
oy["F4"] = "CDR SE by future calendar year"; oy["F4"].font = bold
oy["F5"] = ArrayFormula("F5:O16", f"=PRO.MACK.CDR.YEARS({TRI},,,{ORI})")
oy["A19"] = "ODP bootstrap, 20000 simulations, seed 1"; oy["A19"].font = bold
oy["A20"] = ArrayFormula("A20:H31", f"=PRO.ODP.CDR({TRI},20000,1,,,,{ORI})")
oy["J19"] = "Mack bootstrap, 20000 simulations, seed 1"; oy["J19"].font = bold
oy["J20"] = ArrayFormula("J20:Q31", f"=PRO.MACK.BOOT.CDR({TRI},20000,1,,,,,{ORI})")
oy["A34"] = "Check against R ChainLadder 0.2.21 CDR(MackChainLadder(RAA))"; oy["A34"].font = bold
for h, col in zip(["Quantity", "prospicio", "R", "Match"], "ABCD"):
    oy[f"{col}35"] = h; oy[f"{col}35"].font = bold
for i, (name, f, x) in enumerate([
        ("Total CDR(1) SE", "=C16", 25166.3025264382),
        ("Total Mack SE", "=D16", 26880.7403298914),
        ("Total CDR SE, year 2", "=H16", 7306.52121591382)]):
    row = 36 + i
    oy[f"A{row}"] = name; oy[f"B{row}"] = f; oy[f"C{row}"] = x
    oy[f"D{row}"] = f"=ABS(B{row}-C{row})<1E-6"
oy.column_dimensions["A"].width = 22
for col in "BCDEFGHIJKLMNOPQ":
    oy.column_dimensions[col].width = 12
for row in oy.iter_rows(min_row=6, max_row=38, min_col=2, max_col=17):
    for c in row:
        c.number_format = "#,##0"

# The ODP GLM on the RAA sheet's triangle (Taylor and McGuire, CAS Monograph 3).
og = wb.create_sheet("ODP GLM")
og["A1"] = "prospicio XLL test: the ODP model as a quasi-Poisson GLM on RAA"; og["A1"].font = Font(bold=True, size=13)
og["A2"] = "The reserves are the chain ladder's; RMSEP is the analytic prediction error (delta method)."
og["A2"].font = note
og["A4"] = "Prediction error"; og["A4"].font = bold
og["A5"] = ArrayFormula("A5:H16", f"=PRO.ODP.GLM({TRI},{ORI})")
og["A19"] = "Coefficients and dispersion"; og["A19"].font = bold
og["A20"] = ArrayFormula("A20:D40", f"=PRO.ODP.GLM.PARAMS({TRI},{ORI})")
og["F19"] = "Standardized deviance residuals by development period"; og["F19"].font = bold
og["F20"] = ArrayFormula("F20:I30", f'=PRO.ODP.GLM.RESIDUALS({TRI},"development",{ORI})')
og["F33"] = "Check: total reserve against the chain ladder"; og["F33"].font = bold
og["F34"] = "=D16"; og["G34"] = 52135.2283; og["H34"] = "=ABS(F34-G34)<0.001"
og["K4"] = "Residuals by cell"; og["K4"].font = bold
og["K5"] = ArrayFormula("K5:U60", f"=PRO.ODP.GLM.RESIDUALS({TRI},,{ORI})")
for col in "ABCDEFGHIJKLMNOPQRSTU":
    og.column_dimensions[col].width = 12

# Triangle helpers on the RAA sheet's triangle.
th = wb.create_sheet("Triangle helpers")
th["A1"] = "prospicio XLL test: triangle helpers on RAA"; th["A1"].font = Font(bold=True, size=13)
th["A3"] = "Incremental (PRO.TRI.INCREMENTAL)"; th["A3"].font = bold
th["A4"] = ArrayFormula("A4:J13", f"=PRO.TRI.INCREMENTAL({TRI})")
th["L3"] = "Latest diagonal"; th["L3"].font = bold
th["L4"] = ArrayFormula("L4:L13", f"=PRO.TRI.LATEST({TRI})")
th["A16"] = "Link ratios (PRO.TRI.LINKRATIOS)"; th["A16"].font = bold
th["A17"] = ArrayFormula("A17:I26", f"=PRO.TRI.LINKRATIOS({TRI})")
th["A29"] = "Cumulative again (PRO.TRI.CUMULATIVE of the incremental)"; th["A29"].font = bold
th["A30"] = ArrayFormula("A30:J39", "=PRO.TRI.CUMULATIVE(A4:J13)")
th["L29"] = "Same as RAA"; th["L29"].font = bold
th["L30"] = f"=SUMPRODUCT((A30:J39={TRI})*1)=100"
for row in th.iter_rows(min_row=4, max_row=39, min_col=1, max_col=12):
    for c in row:
        c.number_format = "0.0000" if 17 <= c.row <= 26 else "#,##0"

wb.calculation.fullCalcOnLoad = True

# Exposure curves: a risk profile rated with a Swiss Re curve.
ex = wb.create_sheet("Exposure")
ex["A1"] = "Exposure rating with a Swiss Re curve (MBBEFD)"; ex["A1"].font = bold
ex["A2"] = "Swiss Re c"; ex["B2"] = 3; ex["B2"].fill = inp
ex["A3"] = "Layer limit"; ex["B3"] = 4000000; ex["B3"].fill = inp
ex["A4"] = "Layer attachment"; ex["B4"] = 1000000; ex["B4"].fill = inp
ex["D2"] = ArrayFormula("D2:E6", '=PRO.EXPOSURE.STATS("swiss_re",B2)')
for h, col in zip(["Sum insured", "Premium", "Loss ratio", "Expected loss", "Layer share", "Layer loss"], "ABCDEF"):
    ex[f"{col}8"] = h; ex[f"{col}8"].font = bold
for k, (si, prem) in enumerate([(500000, 2000000), (2000000, 3000000), (5000000, 2500000), (10000000, 1500000), (25000000, 1000000)]):
    r = 9 + k
    ex[f"A{r}"] = si; ex[f"B{r}"] = prem; ex[f"C{r}"] = 0.6
    for c in "ABC":
        ex[f"{c}{r}"].fill = inp
    ex[f"D{r}"] = f"=B{r}*C{r}"
for c in "ABDF":
    for r in range(9, 14):
        ex[f"{c}{r}"].number_format = "#,##0"
ex["E9"] = ArrayFormula("E9:E13", '=PRO.EXPOSURE.SHARE($B$3,$B$4,A9:A13,"swiss_re",$B$2)')
for r in range(9, 14):
    ex[f"E{r}"].number_format = "0.0000"
    ex[f"F{r}"] = f"=D{r}*E{r}"
ex["A14"] = "Total"; ex["A14"].font = bold
ex["D14"] = "=SUM(D9:D13)"; ex["F14"] = "=SUM(F9:F13)"
ex["D14"].number_format = ex["F14"].number_format = "#,##0"
ex["A16"] = "G(x) for c = 1.5, 2, 3, 4, 5"; ex["A16"].font = bold
ex["A17"] = "x"
for k, c in enumerate([1.5, 2, 3, 4, 5]):
    ex.cell(row=17, column=2 + k, value=c).fill = inp
for k, x in enumerate([0, 0.05, 0.1, 0.2, 0.3, 0.5, 0.7, 0.9, 1]):
    ex.cell(row=18 + k, column=1, value=x).fill = inp
ex["B18"] = ArrayFormula("B18:F26", '=PRO.EXPOSURE.CURVE(A18:A26,"swiss_re",B17:F17)')
ex["A28"] = "The c whose mean rate is"; ex["B28"] = 0.05; ex["B28"].fill = inp
ex["C28"] = ArrayFormula("C28", "=PRO.SWISSRE.C(B28)")
ex["A30"] = "The same profile by PRO.PROFILE, net of a surplus of 4 lines on 1,000,000"; ex["A30"].font = bold
ex["A31"] = "Retention"; ex["B31"] = 1000000; ex["B31"].fill = inp
ex["A32"] = "Lines"; ex["B32"] = 4; ex["B32"].fill = inp
ex["A33"] = ArrayFormula("A33:G39", '=PRO.PROFILE(A9:A13,D9:D13,B3,B4,"swiss_re",B2,,,,,,B31,B32)')
for row in ex.iter_rows(min_row=34, max_row=39, min_col=1, max_col=6):
    for c in row:
        c.number_format = "#,##0.00" if c.column == 4 else "#,##0"
for r in range(34, 40):
    ex[f"G{r}"].number_format = "0.0000"
ex.column_dimensions["A"].width = 26
for col in "BCDEFG":
    ex.column_dimensions[col].width = 15


# Credibility: Hachemeister's data (R actuar), and the limited fluctuation table.
cr = wb.create_sheet("Credibility")
cr["A1"] = "Buhlmann-Straub on Hachemeister's data (R actuar): ratios, then weights"; cr["A1"].font = bold
HACHEMEISTER = Path(__file__).resolve().parents[3] / "validation" / "reference" / "hachemeister.csv"
with open(HACHEMEISTER, newline="") as f:
    hach = list(csv.DictReader(f))
for q in range(12):
    cr.cell(row=2, column=2 + q, value=f"Q{q + 1}").font = bold
    cr.cell(row=9, column=2 + q, value=f"Q{q + 1}").font = bold
for st in range(5):
    cr.cell(row=3 + st, column=1, value=f"State {st + 1}")
    cr.cell(row=10 + st, column=1, value=f"State {st + 1}")
for h in hach:
    i, j = int(h["state"]) - 1, int(h["quarter"]) - 1
    cr.cell(row=3 + i, column=2 + j, value=float(h["ratio"])).fill = inp
    cr.cell(row=10 + i, column=2 + j, value=float(h["weight"])).fill = inp
cr["A17"] = ArrayFormula("A17:E22", "=PRO.BUHLMANN(B3:M7,B10:M14,,A3:A7)")
cr["G17"] = ArrayFormula("G17:H21", "=PRO.BUHLMANN.PARAMS(B3:M7,B10:M14)")
cr["J16"] = "Iterative (Bichsel-Straub)"; cr["J16"].font = bold
cr["J17"] = ArrayFormula("J17:K21", '=PRO.BUHLMANN.PARAMS(B3:M7,B10:M14,"iterative")')
cr["A23"] = "Matches R: premium of state 1 is 2055.17"
cr["A24"] = "=ROUND(E18,2)=2055.17"
cr["A26"] = "Limited fluctuation: full credibility standards in claims"; cr["A26"].font = bold
cr["A27"] = "p by k"
for c, k in enumerate([0.025, 0.05, 0.1]):
    cr.cell(row=27, column=2 + c, value=k).fill = inp
for r, p in enumerate([0.9, 0.95, 0.99]):
    cr.cell(row=28 + r, column=1, value=p).fill = inp
cr["B28"] = ArrayFormula("B28:D30", "=PRO.CRED.STANDARD(A28:A30,B27:D27)")
cr["A32"] = "Credibility of 500 claims against each standard"
cr["B32"] = ArrayFormula("B32:D34", "=PRO.CRED.Z(500,B28:D30)")
for row in cr.iter_rows(min_row=28, max_row=30, min_col=2, max_col=4):
    for c in row:
        c.number_format = "#,##0"
for row in cr.iter_rows(min_row=32, max_row=34, min_col=2, max_col=4):
    for c in row:
        c.number_format = "0.000"
cr.column_dimensions["A"].width = 18

# Ratemaking: Werner and Modlin's chapter 5 parallelogram and Table 6.14's trends.
rm = wb.create_sheet("Ratemaking")
rm["A1"] = "Rate level by the parallelogram method (Werner and Modlin, ch. 5)"; rm["A1"].font = bold
rm["A2"] = "Effective"; rm["B2"] = "Change"; rm["A2"].font = rm["B2"].font = bold
for k, (d, c) in enumerate([(datetime.date(2010, 7, 1), 0.05), (datetime.date(2011, 1, 1), 0.10), (datetime.date(2012, 4, 1), -0.01)]):
    rm.cell(row=3 + k, column=1, value=d).number_format = "yyyy-mm-dd"
    rm.cell(row=3 + k, column=2, value=c).number_format = "0.0%"
    rm.cell(row=3 + k, column=1).fill = rm.cell(row=3 + k, column=2).fill = inp
rm["D2"] = "Calendar year"; rm["E2"] = "On-level factor"; rm["F2"] = "Six-month policies"
for h in ("D2", "E2", "F2"):
    rm[h].font = bold
for k, y in enumerate(range(2010, 2014)):
    rm.cell(row=3 + k, column=4, value=y).fill = inp
rm["E3"] = ArrayFormula("E3:E6", "=PRO.ONLEVEL(A3:A5,B3:B5,D3:D6,D3:D6+1)")
rm["F3"] = ArrayFormula("F3:F6", '=PRO.ONLEVEL(A3:A5,B3:B5,D3:D6,D3:D6+1,"calendar",0.5)')
for r in range(3, 7):
    rm[f"E{r}"].number_format = rm[f"F{r}"].number_format = "0.0000"
rm["A9"] = "Calendar year 2011's cells (Figure 5.17)"; rm["A9"].font = bold
rm["A10"] = ArrayFormula("A10:D16", "=PRO.RATELEVEL(A3:A5,B3:B5,2011,2012)")
rm["A18"] = "Trend length: CY2011 losses to rates effective 1/1/2015 for a year"
rm["A19"] = ArrayFormula("A19", '=PRO.AVGDATE(2015,2016,"policy")-PRO.AVGDATE(2011,2012)')
rm["A21"] = "Exponential trend of Table 6.14's severities, quarterly"; rm["A21"].font = bold
sev = [1061.45, 1076.56, 1085.56, 1098.02, 1102.46, 1107.52, 1117.71, 1123.70, 1127.63, 1124.01,
       1116.42, 1109.48, 1105.38, 1102.61, 1096.67, 1093.77, 1100.67, 1108.79, 1122.13, 1126.46]
rm["A22"] = "Time"; rm["B22"] = "Severity"; rm["A22"].font = rm["B22"].font = bold
for k, v in enumerate(sev):
    rm.cell(row=23 + k, column=1, value=2010 + k / 4)
    rm.cell(row=23 + k, column=2, value=v).fill = inp
rm["D22"] = "20 points"; rm["G22"] = "Last 8"; rm["D22"].font = rm["G22"].font = bold
rm["D23"] = ArrayFormula("D23:E28", "=PRO.TREND.FIT(A23:A42,B23:B42)")
rm["G23"] = ArrayFormula("G23:H28", "=PRO.TREND.FIT(A35:A42,B35:B42)")
rm["E24"].number_format = rm["H24"].number_format = "0.0%"
rm["A45"] = "Loss ratio indication (Werner and Modlin, Appendix A)"; rm["A45"].font = bold
heads = ["Earned premium", "On-level", "Premium trend", "Losses", "Development", "Loss trend"]
appendix_a = [
    (1122372, 1.2161, 1.1342, 856495, 1.0000, 0.9912), (1154508, 1.2176, 1.1116, 867184, 0.9799, 0.9962),
    (1280545, 1.1311, 1.0879, 835120, 1.0003, 1.0012), (1369976, 1.0892, 1.0663, 821509, 1.0282, 1.0062),
    (1397750, 1.0991, 1.0452, 797866, 1.0966, 1.0113)]
for c, h in enumerate(heads):
    rm.cell(row=46, column=1 + c, value=h).font = bold
for r, row in enumerate(appendix_a):
    for c, x in enumerate(row):
        rm.cell(row=47 + r, column=1 + c, value=x).fill = inp
rm["H46"] = "ULAE factor"; rm["H46"].font = bold; rm["H47"] = 1.143; rm["H47"].fill = inp
rm["A53"] = ArrayFormula("A53:C59", "=PRO.LR.EXPERIENCE(A47:A51,D47:D51,B47:B51,C47:C51,E47:E51,F47:F51,H47)")
for row in rm.iter_rows(min_row=54, max_row=59, min_col=1, max_col=2):
    for c in row:
        c.number_format = "#,##0"
for r in range(54, 60):
    rm[f"C{r}"].number_format = "0.0%"
rm["A61"] = "Expenses (premium-based): ratio, fixed share"; rm["A61"].font = bold
for r, (x, sh) in enumerate([(0.063, 0.75), (0.085, 0.75), (0.002, 1.0), (0.112, 0.0), (0.021, 0.0)]):
    rm.cell(row=62 + r, column=1, value=x).fill = inp
    rm.cell(row=62 + r, column=2, value=sh).fill = inp
rm["D62"] = ArrayFormula("D62:E63", "=PRO.EXPENSES(A62:A66,B62:B66)")
rm["D65"] = "Profit"; rm["E65"] = 0.05; rm["E65"].fill = inp
rm["D66"] = "Indicated change"
rm["E66"] = ArrayFormula("E66", "=PRO.INDICATED.CHANGE(C59,E63,E65,E62)")
rm["D67"] = "Permissible loss ratio"
rm["E67"] = ArrayFormula("E67", "=PRO.PERMISSIBLE.LR(E63,E65,E62)")
for c in ("E62", "E63", "E66", "E67"):
    rm[c].number_format = "0.0%"
rm.column_dimensions["A"].width = 16
for col in "BCDEFGH":
    rm.column_dimensions[col].width = 15

# Burning cost: Clark's per-risk example (Basics of Reinsurance Pricing, A.4.a).
bc = wb.create_sheet("Burning cost")
bc["A1"] = "Burning cost of 400,000 xs 100,000, trended at 4% to 7/1/1997 (Clark, A.4.a)"; bc["A1"].font = bold
bc["A2"] = "Accident date"; bc["B2"] = "Amount"; bc["A2"].font = bc["B2"].font = bold
clark = [((1988, 9, 20), 240946), ((1988, 10, 11), 821499), ((1989, 3, 15), 158129), ((1990, 6, 21), 114051),
         ((1990, 10, 24), 78043), ((1991, 1, 10), 162533), ((1992, 2, 23), 324298), ((1992, 4, 30), 100549),
         ((1992, 9, 22), 75476), ((1993, 1, 1), 171885), ((1993, 5, 18), 94218), ((1993, 8, 1), 170297),
         ((1994, 8, 15), 87133), ((1995, 7, 12), 771249)]
for r, ((y, m, d), x) in enumerate(clark):
    c = bc.cell(row=3 + r, column=1, value=datetime.date(y, m, d)); c.number_format = "yyyy-mm-dd"; c.fill = inp
    c = bc.cell(row=3 + r, column=2, value=x); c.number_format = "#,##0"; c.fill = inp
bc["D2"] = "Year"; bc["E2"] = "Premium"; bc["F2"] = "Excess LDF"
for h in ("D2", "E2", "F2"):
    bc[h].font = bold
years = [(1988, 1422554, 1.0), (1989, 1823103, 1.0), (1990, 2054034, 1.0), (1991, 2147147, 1.0),
         (1992, 2151541, 1.01), (1993, 2159198, 1.05), (1994, 2167158, 1.15), (1995, 2187654, 1.3)]
for r, (y, p_, f) in enumerate(years):
    for c, x in enumerate((y, p_, f)):
        bc.cell(row=3 + r, column=4 + c, value=x).fill = inp
    bc.cell(row=3 + r, column=5).number_format = "#,##0"
bc["H2"] = "Limit"; bc["H3"] = 400000; bc["I2"] = "Attachment"; bc["I3"] = 100000
bc["H4"] = "Trend"; bc["H5"] = 0.04; bc["I4"] = "Trended to"; bc["I5"] = datetime.date(1997, 7, 1)
bc["I5"].number_format = "yyyy-mm-dd"
for c in ("H3", "I3", "H5", "I5"):
    bc[c].fill = inp
bc["A19"] = ArrayFormula("A19:G29", "=PRO.BURNINGCOST(A3:A16,B3:B16,H3,I3,D3:D10,E3:E10,F3:F10,H5,I5)")
for row in bc.iter_rows(min_row=20, max_row=29, min_col=2, max_col=6):
    for c in row:
        c.number_format = "0.000" if c.column == 4 else "#,##0"
for r in range(20, 30):
    bc[f"G{r}"].number_format = "0.0%"
bc.column_dimensions["A"].width = 14
for col in "BCDEFGHI":
    bc.column_dimensions[col].width = 16

# Aggregate distributions: a compound total and a layer table.
ag = wb.create_sheet("Aggregate")
ag["A1"] = "Annual aggregate loss: 3 claims a year (dispersion 2), gamma severity mean 50,000, CV 1.5"; ag["A1"].font = bold
for r, (label, x) in enumerate([("Claims a year", 3), ("Dispersion", 2), ("Family", "gamma_cv"), ("Mean", 50000), ("CV", 1.5)]):
    ag.cell(row=2 + r, column=1, value=label)
    ag.cell(row=2 + r, column=2, value=x).fill = inp
ag["D1"] = "Ground up"; ag["G1"] = "100,000 xs 50,000 per claim"; ag["D1"].font = ag["G1"].font = bold
ag["D2"] = ArrayFormula("D2:E18", "=PRO.AGG.STATS(B2,B3,B4,B5,B6)")
ag["G2"] = ArrayFormula("G2:H18", "=PRO.AGG.STATS(B2,B3,B4,B5,B6,,100000,50000)")
for c in ("E", "H"):
    for r in range(2, 19):
        ag[f"{c}{r}"].number_format = "#,##0.0000" if r in (4, 5) or r >= 16 else "#,##0"
ag["A21"] = "Per-occurrence layers (aggregate's reins_stats_df)"; ag["A21"].font = bold
ag["A22"] = "Limits"; ag["A23"] = "Attachments"
for c, (l, a) in enumerate([(50000, 50000), (100000, 100000), (300000, 200000)]):
    ag.cell(row=22, column=2 + c, value=l).fill = inp
    ag.cell(row=23, column=2 + c, value=a).fill = inp
ag["A25"] = ArrayFormula("A25:D39", "=PRO.AGG.LAYERS(B2,B3,B4,B5,B6,,B22:D22,B23:D23)")
for row in ag.iter_rows(min_row=26, max_row=39, min_col=2, max_col=4):
    for c in row:
        c.number_format = "0.0000"
ag["A41"] = "Occurrence and aggregate PML curves (aggregate's oep beside the total's VaR)"; ag["A41"].font = bold
ag["A42"] = ArrayFormula("A42:F51", "=PRO.AGG.PML(B2,B3,B4,B5,B6)")
for row in ag.iter_rows(min_row=43, max_row=51, min_col=1, max_col=6):
    for c in row:
        c.number_format = "0.0000" if c.column in (2, 5) else "#,##0.0" if c.column == 6 else "#,##0"
ag["A54"] = "Stop loss: 150,000 xs 200,000 a year, net of 100,000 xs 50,000 per claim (aggregate's agg block)"; ag["A54"].font = bold
ag["A55"] = ArrayFormula("A55:B63", "=PRO.AGG.COVER(B2,B3,B4,B5,B6,,150000,200000,,100000,50000)")
for r in range(56, 64):
    ag[f"B{r}"].number_format = "0.0000" if r in (57, 58, 61, 62, 63) else "#,##0"
ag["A66"] = "Tower priced on the grid by a proportional hazard of 0.7: two layers and a stop loss on the net"; ag["A66"].font = bold
ag["A67"] = "Limits"; ag["A68"] = "Attachments"
for c, (l, a) in enumerate([(100000, 50000), (300000, 150000)]):
    ag.cell(row=67, column=2 + c, value=l).fill = inp
    ag.cell(row=68, column=2 + c, value=a).fill = inp
ag["A70"] = ArrayFormula("A70:H75", '=PRO.TOWER.PRICE(B2,B3,B4,B5,B6,,,,B67:C67,B68:C68,,,,200000,400000,"ph",0.7)')
for row in ag.iter_rows(min_row=71, max_row=75, min_col=2, max_col=8):
    for c in row:
        c.number_format = "0.0000" if c.column in (4, 7, 8) else "#,##0"
ag["J66"] = "Contract terms on 100,000 xs 50,000 in expectation: deposit 60,000, one reinstatement at 100%"; ag["J66"].font = bold
ag["J67"] = ArrayFormula("J67:K74", "=PRO.LAYER.TERMS(B2,B3,B4,B5,B6,,100000,50000,60000,1)")
ag["M66"] = "With a sliding scale (35% at 50% down to 25% at 70%) and a 20% profit commission after 30%"; ag["M66"].font = bold
ag["M67"] = ArrayFormula("M67:N74", "=PRO.LAYER.TERMS(B2,B3,B4,B5,B6,,100000,50000,60000,,{0.35,0.25},{0.5,0.7},0.2,0.3)")
ag.column_dimensions["A"].width = 18
ag.column_dimensions["D"].width = 16
ag.column_dimensions["G"].width = 16
for col in "BCEH":
    ag.column_dimensions[col].width = 14

# Severity fits to large losses: R Pareto's example losses, some reported
# only above 1,000 and three capped by a policy limit.
sf = wb.create_sheet("Severity fit")
sf["A1"] = "Large losses fitted by maximum likelihood, each conditioned on its reporting threshold; TRUE where a policy limit capped it"; sf["A1"].font = bold
sf["A3"] = "Loss"; sf["B3"] = "Reported above"; sf["C3"] = "Capped"
for c in ("A3", "B3", "C3"):
    sf[c].font = bold
losses = [1100, 1300, 1750, 2000, 2600, 3500, 4100, 5200, 7000, 9000, 12000, 18000, 25000, 40000]
reported = [1000, 1200, 1000, 1500, 1000, 1000, 3000, 1000, 1000, 5000, 1000, 1000, 1000, 1000]
capped = [False] * 5 + [True] + [False] * 4 + [True, False, False, True]
for i, (y, r, c) in enumerate(zip(losses, reported, capped)):
    for col, v in zip("ABC", (y, r, c)):
        sf[f"{col}{4 + i}"] = v
        sf[f"{col}{4 + i}"].fill = inp
data = "A4:A17,{fam},{out},{t},,B4:B17,C4:C17"
sf["E3"] = "Lognormal"; sf["E3"].font = bold
sf["E4"] = ArrayFormula("E4:G13", "=PRO.SEV.FIT(" + data.format(fam='"lognormal"', out="", t="") + ")")
sf["I3"] = "Pareto above 1,000"; sf["I3"].font = bold
sf["I4"] = ArrayFormula("I4:K13", "=PRO.SEV.FIT(" + data.format(fam='"pareto"', out="", t="1000") + ")")
sf["E14"] = "The lognormal's 1M xs 1M from the fit's cells"
sf["I14"] = ArrayFormula("I14", "=PRO.LAYER(1000000,1000000,\"lognormal\",F5,F6)")
sf["E16"] = "Body families compared"; sf["E16"].font = bold
sf["E17"] = ArrayFormula("E17:N22", "=PRO.SEV.FIT(" + data.format(fam="", out='"compare"', t="") + ")")
sf["E23"] = "Lognormal against the product-limit estimate (P-P and Q-Q)"; sf["E23"].font = bold
sf["E24"] = ArrayFormula("E24:J35", "=PRO.SEV.FIT(" + data.format(fam='"lognormal"', out='"gof"', t="") + ")")
sf["P3"] = "Grouped losses (Bahnemann's six bands; the last bound blank: open)"; sf["P3"].font = bold
sf["P4"] = "From"; sf["Q4"] = "Count"
for c in ("P4", "Q4"):
    sf[c].font = bold
for i, (b, c) in enumerate(zip([0, 1000, 1500, 2000, 2500, 3000, None], [42, 61, 47, 26, 14, 10, None])):
    sf[f"P{5 + i}"] = b; sf[f"P{5 + i}"].fill = inp
    if c is not None:
        sf[f"Q{5 + i}"] = c; sf[f"Q{5 + i}"].fill = inp
sf["S3"] = "Gamma by minimum chi-square (Bahnemann P2.16: 4.7432, 337.31)"; sf["S3"].font = bold
sf["S4"] = ArrayFormula("S4:U12", '=PRO.SEV.FIT.GROUPED(P5:P11,Q5:Q10,"gamma","chi_square")')
sf["S14"] = "Observed and expected per band"; sf["S14"].font = bold
sf["S15"] = ArrayFormula("S15:V21", '=PRO.SEV.FIT.GROUPED(P5:P11,Q5:Q10,"gamma","chi_square","table")')
sf["P24"] = "Limited means (Bahnemann Ex 2.10: 1,500 claims under a 300,000 limit)"; sf["P24"].font = bold
sf["P25"] = "Size"; sf["Q25"] = "E[X; x]"
for c in ("P25", "Q25"):
    sf[c].font = bold
x = [1000, 5000, 10000, 25000, 50000, 100000, 200000, 300000]
below = [729, 1096, 1208, 1326, 1391, 1440, 1468, 1500]
total = [225138, 1102272, 1918947, 3752091, 6007543, 9234739, 13100561, 22343455]
for i, (v, b, t) in enumerate(zip(x, below, total)):
    sf[f"P{26 + i}"] = v; sf[f"P{26 + i}"].fill = inp
    sf[f"Q{26 + i}"] = f"=({t}+P{26 + i}*(1500-{b}))/1500"
sf["S24"] = "Lognormal by minimum distance (6.9852, 2.5850)"; sf["S24"].font = bold
sf["S25"] = ArrayFormula("S25:T28", '=PRO.SEV.FIT.LEV(P26:P33,Q26:Q33,"lognormal")')
sf["S30"] = ArrayFormula("S30:U38", '=PRO.SEV.FIT.LEV(P26:P33,Q26:Q33,"lognormal",,"table")')
for col in "ABCEFGHIJKLMPQSTUV":
    sf.column_dimensions[col].width = 14

# Extreme value theory on 1,000 exact Burr quantiles (R ismev's test data).
ev = wb.create_sheet("EVT")
ev["A1"] = "1,000 large losses (exact quantiles of a Burr with a = 2, c = 1.5): choose a threshold, fit a GPD above it"; ev["A1"].font = bold
ev["A3"] = "Losses"; ev["A3"].font = bold
for i in range(1000):
    p = (i + 0.5) / 1000
    ev[f"A{4 + i}"] = ((1 - p) ** (-1 / 2) - 1) ** (1 / 1.5)
ev["C3"] = "Threshold diagnostics: where the shape and modified scale stop moving, a GPD fits"; ev["C3"].font = bold
ev["C4"] = ArrayFormula("C4:K14", "=PRO.EVT.THRESHOLDS(A4:A1003)")
ev["C17"] = "Threshold"; ev["D17"] = 1; ev["D17"].fill = inp
ev["C18"] = ArrayFormula("C18:E29", "=PRO.EVT.FIT(A4:A1003,D17)")
ev["G17"] = "Hill estimates"; ev["G17"].font = bold
ev["G18"] = ArrayFormula("G18:J28", "=PRO.EVT.HILL(A4:A1003)")
for col in "CDEFGHIJK":
    ev.column_dimensions[col].width = 14

# The time value of money: RAA's paid chain ladder as a payment pattern.
pv = wb.create_sheet("Payout")
pv["A1"] = "RAA's factors to ultimate (the RAA sheet's CDF row) as a payment pattern, losses occurring evenly over the year"; pv["A1"].font = bold
pv["A3"] = "Years"; pv["A4"] = "CDF"; pv["A3"].font = pv["A4"].font = bold
for c in range(10):
    pv.cell(row=3, column=2 + c, value=c + 1).fill = inp
    pv.cell(row=4, column=2 + c, value=f"=RAA!{chr(66 + c)}15")
pv["A6"] = "Discount rate"; pv["B6"] = 0.04; pv["B6"].fill = inp
pv["A7"] = "Cost of capital"; pv["B7"] = 0.10; pv["B7"].fill = inp
pv["A8"] = "Expected loss"; pv["B8"] = 1000000; pv["B8"].fill = inp
pv["A9"] = "Assets (one-year VaR)"; pv["B9"] = 2500000; pv["B9"].fill = inp
pv["A11"] = "Pattern"; pv["A11"].font = bold
pv["A12"] = ArrayFormula("A12:E22", '=PRO.PAYOUT.PATTERN(B3:K3,B4:K4,"factors",0.5,B6)')
pv["G11"] = "Statistics"; pv["G11"].font = bold
pv["G12"] = ArrayFormula("G12:H14", '=PRO.PAYOUT.STATS(B3:K3,B4:K4,B6,"factors",0.5)')
pv["G16"] = "Cost-of-capital price over the payout"; pv["G16"].font = bold
pv["G17"] = ArrayFormula("G17:H25", '=PRO.PAYOUT.PRICE(B8,B9,B7,B3:K3,B4:K4,B6,"factors",0.5)')
for col in "ABCDEFGH":
    pv.column_dimensions[col].width = 16
pv["A27"] = "A cover's quote: the guide's middle layer, risk premium 1.24 on an expected loss of 0.586, limit 5"; pv["A27"].font = bold
pv["A28"] = "Expected loss"; pv["B28"] = 0.586; pv["B28"].fill = inp
pv["A29"] = "Risk premium"; pv["B29"] = 1.24; pv["B29"].fill = inp
pv["A30"] = "Assets"; pv["B30"] = 7.778; pv["B30"].fill = inp
pv["A31"] = "Expenses"; pv["B31"] = 0.03; pv["B31"].fill = inp
pv["A32"] = "Brokerage"; pv["B32"] = 0.10; pv["B32"].fill = inp
pv["A33"] = "Limit"; pv["B33"] = 5; pv["B33"].fill = inp
pv["A34"] = "Quoted"; pv["B34"] = 1.30; pv["C34"] = 1.60; pv["B34"].fill = pv["C34"].fill = inp
pv["E27"] = ArrayFormula("E27:H40", "=PRO.QUOTE(B28,B29,B30,B31,B32,,B33,B34:C34)")

# Distortions: Mildenhall and Major's InsCo, priced and calibrated.
ds = wb.create_sheet("Distortions")
ds["A1"] = "InsCo (Pricing Insurance Risk): seven totals, assets 100; every standard distortion at the 15% CCoC premium"; ds["A1"].font = bold
ds["A3"] = "Total"; ds["B3"] = "Probability"; ds["A3"].font = ds["B3"].font = bold
for i, (x, p) in enumerate(zip([22, 28, 36, 40, 55, 65, 100], [0.1, 0.1, 0.1, 0.4, 0.1, 0.1, 0.1])):
    ds[f"A{4 + i}"] = x; ds[f"B{4 + i}"] = p
    ds[f"A{4 + i}"].fill = ds[f"B{4 + i}"].fill = inp
ds["A12"] = "Target premium"; ds["B12"] = "=(46.6+0.15*100)/1.15"
ds["D3"] = "Calibrated to the target"; ds["D3"].font = bold
ds["D4"] = ArrayFormula("D4:F9", "=PRO.DISTORTION.CALIBRATE(A4:A10,B4:B10,B12)")
ds["D12"] = "The PH at its calibrated rho, priced back"; ds["D12"].font = bold
ds["D13"] = ArrayFormula("D13:E16", '=PRO.DISTORTION.PRICE(A4:A10,B4:B10,"ph",E6)')
ds["H3"] = "s"; ds["I3"] = "PH"; ds["J3"] = "Wang"; ds["K3"] = "TVaR"
for c in ("H3", "I3", "J3", "K3"):
    ds[c].font = bold
for i in range(11):
    ds[f"H{4 + i}"] = i / 10
ds["I4"] = ArrayFormula("I4:I14", '=PRO.DISTORTION(H4:H14,"ph",E6)')
ds["J4"] = ArrayFormula("J4:J14", '=PRO.DISTORTION(H4:H14,"wang",E7)')
ds["K4"] = ArrayFormula("K4:K14", '=PRO.DISTORTION(H4:H14,"tvar",E9)')
ds["L3"] = "Bi-TVaR"; ds["M3"] = "Beta"; ds["L3"].font = ds["M3"].font = bold
ds["O3"] = "Bi-TVaR p0, p1, w"; ds["O3"].font = bold
for i, v in enumerate([0.2, 0.9, 0.3]):
    ds.cell(row=4, column=15 + i, value=v).fill = inp
ds["O6"] = "Beta a, b"; ds["O6"].font = bold
for i, v in enumerate([0.6, 1.5]):
    ds.cell(row=7, column=15 + i, value=v).fill = inp
ds["L4"] = ArrayFormula("L4:L14", '=PRO.DISTORTION(H4:H14,"bitvar",O4:Q4)')
ds["M4"] = ArrayFormula("M4:M14", '=PRO.DISTORTION(H4:H14,"beta",O7:P7)')
ds["O9"] = "Weighted TVaR: levels | weights"; ds["O9"].font = bold
for i, (p, w) in enumerate([(0.1, 0.2), (0.5, 0.5), (0.9, 0.3)]):
    ds[f"O{10 + i}"] = p; ds[f"P{10 + i}"] = w
    ds[f"O{10 + i}"].fill = ds[f"P{10 + i}"].fill = inp
ds["D18"] = "InsCo priced by the weighted TVaR"; ds["D18"].font = bold
ds["D19"] = ArrayFormula("D19:E22", '=PRO.DISTORTION.PRICE(A4:A10,B4:B10,"wtdtvar",O10:P12)')
for col in "ABDEFHIJKLMOPQ":
    ds.column_dimensions[col].width = 14

# Copulas: two lines joined by a Clayton copula, and the same draws reordered by Iman-Conover.
cp = wb.create_sheet("Copulas")
cp["A1"] = "Two lines (lognormal, mean 50,000, CV 3 and gamma, mean 50,000, CV 1) joined by a Clayton copula; 1,000 simulated years"; cp["A1"].font = bold
cp["A3"] = "Clayton theta"; cp["B3"] = 2; cp["B3"].fill = inp
cp["A4"] = "Seed"; cp["B4"] = 1; cp["B4"].fill = inp
cp["A5"] = "Target rank correlation"; cp["B5"] = 0.5; cp["B5"].fill = inp
cp["A7"] = "C(u, v) at"; cp["A7"].font = bold
for i, (u, v) in enumerate([(0.1, 0.1), (0.5, 0.5), (0.9, 0.9), (0.4, 0.6)]):
    cp[f"A{8 + i}"] = u; cp[f"B{8 + i}"] = v
    cp[f"A{8 + i}"].fill = cp[f"B{8 + i}"].fill = inp
cp["C7"] = "Clayton"; cp["D7"] = "Gaussian at 0.5"; cp["C7"].font = cp["D7"].font = bold
cp["C8"] = ArrayFormula("C8:C11", '=PRO.COPULA.CDF(A8:B11,"clayton",B3)')
cp["D8"] = ArrayFormula("D8:D11", '=PRO.COPULA.CDF(A8:B11,"gaussian",0.5)')
cp["F3"] = "U"; cp["G3"] = "V"; cp["H3"] = "Line A"; cp["I3"] = "Line B"; cp["J3"] = "Total"
cp["L3"] = "A, reordered"; cp["M3"] = "B, reordered"
for c in ("F3", "G3", "H3", "I3", "J3", "L3", "M3"):
    cp[c].font = bold
cp["F4"] = ArrayFormula("F4:G1003", '=PRO.COPULA.SIM("clayton",B3,1000,2,B4)')
cp["H4"] = ArrayFormula("H4:H1003", '=PRO.SEV.QUANTILE(F4:F1003,"lognormal_cv",50000,3)')
cp["I4"] = ArrayFormula("I4:I1003", '=PRO.SEV.QUANTILE(G4:G1003,"gamma_cv",50000,1)')
for r in range(4, 1004):
    cp[f"J{r}"] = f"=H{r}+I{r}"
cp["L4"] = ArrayFormula("L4:M1003", "=PRO.IMAN.CONOVER(H4:I1003,B5,B4)")
cp["A14"] = "Total's 99th percentile"; cp["B14"] = "=PERCENTILE(J4:J1003,0.99)"
for col in "ABCDFGHIJLM":
    cp.column_dimensions[col].width = 14

# A three-year term: profit commission deficits and a sliding scale's deficits and credits carried forward.
tm = wb.create_sheet("Term")
tm["A1"] = "Three-year terms: one row a path, one column a year; premium 100 a year at 100%"; tm["A1"].font = bold
tm["A3"] = "Premium"; tm["B3"] = 100; tm["B3"].fill = inp
tm["A4"] = "Profit share"; tm["B4"] = 0.5; tm["B4"].fill = inp
tm["A5"] = "Allowance"; tm["B5"] = 0.2; tm["B5"].fill = inp
tm["A6"] = "Carry years (blank: until used up)"; tm["B6"].fill = inp
tm["A7"] = "Commissions"; tm["B7"] = 0.4; tm["C7"] = 0.2
tm["A8"] = "At loss ratios"; tm["B8"] = 0.5; tm["C8"] = 0.7
for c in ("B7", "C7", "B8", "C8"):
    tm[c].fill = inp
tm["A10"] = "Ceded losses"; tm["E10"] = "Profit commission"; tm["I10"] = "Sliding-scale commission"
for c in ("A10", "E10", "I10"):
    tm[c].font = bold
for i, path in enumerate([(120, 70, 30), (90, 40, 30), (60, 60, 60), (30, 110, 50), (85, 50, 95)]):
    for j, x in enumerate(path):
        cell = tm.cell(row=11 + i, column=1 + j, value=x)
        cell.fill = inp
tm["E11"] = ArrayFormula("E11:G15", "=PRO.PROFIT.COMMISSION.TERM(A11:C15,B3,B4,B5,B6)")
tm["I11"] = ArrayFormula("I11:K15", '=PRO.SLIDING.SCALE.TERM(A11:C15,B3,B7:C7,B8:C8,"both",B6)')
tm["A17"] = "Year by year without a carry, for comparison"; tm["A17"].font = bold
tm["E17"] = ArrayFormula("E17:G21", "=PRO.PROFIT.COMMISSION(A11:C15,B3,B4,B5)")
tm["I17"] = ArrayFormula("I17:K21", "=PRO.SLIDING.SCALE(A11:C15,B3,B7:C7,B8:C8)")
tm.column_dimensions["A"].width = 30

# Bayesian MCMC: the ODP chain ladder sampled with NUTS on RAA (England and Verrall).
by = wb.create_sheet("Bayesian")
by["A1"] = "Bayesian ODP chain ladder on RAA, sampled with NUTS (4 chains, 1,000 warm-up and 1,000 kept draws each)"; by["A1"].font = bold
by["A2"] = "The triangle is the RAA sheet's; change the seed or draws below and the fit reruns."
by["A3"] = "Draws per chain"; by["B3"] = 1000; by["B3"].fill = inp
by["A4"] = "Chains"; by["B4"] = 4; by["B4"].fill = inp
by["A5"] = "Seed"; by["B5"] = 0; by["B5"].fill = inp
by["A6"] = "Prior SD"; by["B6"] = 10; by["B6"].fill = inp
args = "RAA!B4:K13,{out},B3,B4,B5,,,B6"
by["A8"] = "Diagnostics (check these first: R-hat under 1.01, ESS above 400, no divergences)"; by["A8"].font = bold
by["A9"] = ArrayFormula("A9:B16", "=PRO.BAYES.ODP(" + args.format(out='"diagnostics"') + ")")
by["D8"] = "Reserves: chain ladder, posterior mean, SD, CV and percentiles"; by["D8"].font = bold
by["D9"] = ArrayFormula("D9:L20", "=PRO.BAYES.ODP(" + args.format(out='"reserves"') + ")")
by["A23"] = "Posterior of the log-scale parameters: c, the origin effects a[w] and development effects b[d]"; by["A23"].font = bold
by["A24"] = ArrayFormula("A24:I43", "=PRO.BAYES.ODP(" + args.format(out='"params"') + ")")
by["D22"] = "Compare: the ODP bootstrap's total SD"
by["H22"] = ArrayFormula("H22", "=INDEX(PRO.ODP(RAA!B4:K13,20000),12,4)")
for row in by.iter_rows(min_row=10, max_row=20, min_col=5, max_col=12):
    for c in row:
        c.number_format = "0.000" if c.column == 8 else "#,##0"
for row in by.iter_rows(min_row=25, max_row=43, min_col=2, max_col=9):
    for c in row:
        c.number_format = "0.0000" if c.column <= 7 else "#,##0"
by.column_dimensions["A"].width = 20
for col in "BCDEFGHIJKL":
    by.column_dimensions[col].width = 13
# The Bayesian Bornhuetter-Ferguson on the Expected loss sheet's GenIns and premium:
# a gamma prior on each ultimate at the a priori LR times premium (decision 38).
EG, EP, EO = "'Expected loss'!$B$4:$K$13", "'Expected loss'!$L$4:$L$13", "'Expected loss'!$A$4:$A$13"
by["A46"] = "Bayesian Bornhuetter-Ferguson on GenIns (Expected loss sheet), prior CV below"; by["A46"].font = bold
by["A47"] = "Prior CV"; by["B47"] = 0.1; by["B47"].fill = inp
bfargs = EG + "," + EP + ",'Expected loss'!$O$3,$B$47,{out},B3,B4,B5,,,,," + EO
by["A49"] = "Diagnostics"; by["A49"].font = bold
by["A50"] = ArrayFormula("A50:B57", "=PRO.BAYES.BF(" + bfargs.format(out='"diagnostics"') + ")")
by["D49"] = "Prior and posterior ultimates"; by["D49"].font = bold
by["D50"] = ArrayFormula("D50:I61", "=PRO.BAYES.BF(" + bfargs.format(out='"ultimates"') + ")")
by["K49"] = "Prior total = LR x premium"
by["K50"] = "=ABS(INDEX(D50:I61,12,4)-'Expected loss'!$O$3*SUM(" + EP + "))<1E-6"
by["K51"] = "R-hat under 1.01"; by["K52"] = "=B50<1.01"
by["A63"] = "Reserves: chain ladder, posterior mean, SD, CV and percentiles"; by["A63"].font = bold
by["A64"] = ArrayFormula("A64:I75", "=PRO.BAYES.BF(" + bfargs.format(out='"reserves"') + ")")
for row in by.iter_rows(min_row=51, max_row=61, min_col=5, max_col=9):
    for c in row:
        c.number_format = "0.000" if c.column == 9 else "#,##0"
for row in by.iter_rows(min_row=65, max_row=75, min_col=2, max_col=9):
    for c in row:
        c.number_format = "0.000" if c.column == 5 else "#,##0"
# The twelve functions the sheets above leave out, each beside a check.
ly["A70"] = "Distribution and survival functions (lognormal mean 50,000, CV 3)"; ly["A70"].font = bold
for h, col in zip(["x", "CDF", "Survival", "Sum is 1"], "ABCD"):
    ly[f"{col}71"] = h; ly[f"{col}71"].font = bold
for k, x in enumerate([10000, 50000, 250000, 1000000]):
    ly.cell(row=72 + k, column=1, value=x).fill = inp
ly["B72"] = ArrayFormula("B72:B75", '=PRO.SEV.CDF(A72:A75,"lognormal_cv",$B$6,$B$7)')
ly["C72"] = ArrayFormula("C72:C75", '=PRO.SEV.SURVIVAL(A72:A75,"lognormal_cv",$B$6,$B$7)')
for r in range(72, 76):
    ly[f"D{r}"] = f"=ABS(B{r}+C{r}-1)<1E-12"
    ly[f"B{r}"].number_format = ly[f"C{r}"].number_format = "0.000000"
ly["A77"] = "Alpha between a frequency and a layer (R Pareto: 3 losses above 800, expected loss 1,200 to 1000 xs 1000)"
ly["G77"] = ArrayFormula("G77", "=PRO.PARETO.ALPHA.FREQLAYER(800,3,1000,1000,1200)")
ly["H77"] = 1.5853501938706767; ly["I77"] = "=ABS(G77-H77)<1E-9"

rm["A70"] = "Pure premium indication (Werner and Modlin, Appendix B)"; rm["A70"].font = bold
for c, h in enumerate(["Exposure", "Losses", "Development", "Loss trend"]):
    rm.cell(row=71, column=1 + c, value=h).font = bold
appendix_b = [(12760, 5161624, 1.0000, 1.1939), (12766, 4820968, 1.0012, 1.1705), (12805, 4112172, 1.0055, 1.1476),
              (12834, 5052052, 1.0186, 1.1251), (13411, 6559224, 1.0555, 1.1030)]
for r, row in enumerate(appendix_b):
    for c, x in enumerate(row):
        rm.cell(row=72 + r, column=1 + c, value=x).fill = inp
rm["F71"] = "LAE factor"; rm["F71"].font = bold; rm["F72"] = 1.012; rm["F72"].fill = inp
rm["A78"] = ArrayFormula("A78:C84", "=PRO.PP.EXPERIENCE(A72:A76,B72:B76,C72:C76,D72:D76,F72)")
rm["E78"] = "Book: pure premium 469.88"; rm["E79"] = "=ROUND(C84,2)=469.88"
rm["E81"] = "Credibility 0.795 on 469.88, complement 585.75 (book 493.63)"
rm["E82"] = ArrayFormula("E82", "=PRO.CRED.WEIGHTED(469.88,0.795,585.75)")
rm["F82"] = "=ROUND(E82,2)=493.63"
rm["E84"] = "Indicated rate: fixed 77.83, variable 13.8%, profit 5% (book 850.76)"
rm["E85"] = ArrayFormula("E85", "=PRO.INDICATED.RATE(493.63+29.11+74.57+15.68,0.138,0.05,77.83)")
rm["F85"] = "=ROUND(E85,2)=850.76"
rm["A87"] = "Trended present rates complement (Appendix A: 13.2% indicated, 5% taken, trend 1.005/1.02; book 6.2%)"
rm["E88"] = ArrayFormula("E88", "=PRO.TRENDED.PRESENT(0.132,0.05,1.005/1.02)")
rm["F88"] = "=ROUND(E88,3)=0.062"
rm["E88"].number_format = "0.0%"
rm["A90"] = "Excel dates as half-month year fractions"; rm["A90"].font = bold
for k, (d, want) in enumerate([((2013, 6, 30), 2013.5), ((2011, 11, 15), 2011.875), ((2015, 12, 31), 2016.0), ((2010, 7, 1), 2010.5)]):
    c = rm.cell(row=91 + k, column=1, value=datetime.date(*d)); c.number_format = "yyyy-mm-dd"; c.fill = inp
    rm.cell(row=91 + k, column=3, value=want)
    rm.cell(row=91 + k, column=4, value=f"=B{91 + k}=C{91 + k}")
rm["B91"] = ArrayFormula("B91:B94", "=PRO.YEARFRACTION(A91:A94)")

cr["A37"] = "Credibility-weighted estimate: Z 0.30 on 200 against 300 (Werner and Modlin, ch. 12: 270)"
cr["F37"] = ArrayFormula("F37", "=PRO.CRED.WEIGHTED(200,0.3,300)")
cr["G37"] = "=ROUND(F37,2)=270"

ex["A42"] = "Destruction rate quantiles (for simulating losses), one column per c"; ex["A42"].font = bold
ex["A43"] = "u"
for k, c in enumerate([1.5, 2, 3, 4, 5]):
    ex.cell(row=43, column=2 + k, value=c).fill = inp
for k, u in enumerate([0.5, 0.9, 0.99, 0.999]):
    ex.cell(row=44 + k, column=1, value=u).fill = inp
ex["B44"] = ArrayFormula("B44:F47", '=PRO.EXPOSURE.QUANTILE(A44:A47,"swiss_re",B43:F43)')
ex["A49"] = "The c whose share of 2.5m xs 2.5m on a 10m risk is that of c = 3"
ex["B49"] = ArrayFormula("B49", '=PRO.EXPOSURE.SHARE(2500000,2500000,10000000,"swiss_re",3)')
ex["C49"] = ArrayFormula("C49", "=PRO.SWISSRE.C.LAYER(2500000,2500000,10000000,B49)")
ex["D49"] = "=ABS(C49-3)<1E-8"
for row in ex.iter_rows(min_row=44, max_row=47, min_col=2, max_col=6):
    for c in row:
        c.number_format = "0.0000"

tm["A24"] = "Swing (retro) premium: basic 0, loss conversion 1, between 100 and 300 (aggregate's swing_a)"; tm["A24"].font = bold
for k, (x, want) in enumerate([(50, 100), (200, 200), (400, 300)]):
    tm.cell(row=25 + k, column=1, value=x).fill = inp
    tm.cell(row=25 + k, column=3, value=want)
    tm.cell(row=25 + k, column=4, value=f"=B{25 + k}=C{25 + k}")
tm["B25"] = ArrayFormula("B25:B27", "=PRO.SWING(A25:A27,0,1,100,300)")

md = wb.create_sheet("Mack draws")
md["A1"] = "prospicio XLL test: Mack's bootstrap draws (PRO.MACK.BOOT.SIMS)"; md["A1"].font = Font(bold=True, size=13)
md["A2"] = "Simulations and seed come from the Mack bootstrap sheet; the average total is that sheet's Mean."
md["A2"].font = note
md["N3"] = "Check"; md["N3"].font = bold
md["N4"] = "Mean total, here"; md["O4"] = "=AVERAGE(K5:K100004)"
md["N5"] = "Mean total, Mack bootstrap sheet"; md["O5"] = "='Mack bootstrap'!C19"
md["N6"] = "Match"; md["O6"] = "=ABS(O4-O5)<1E-6*O5"
md["A4"] = ArrayFormula("A4:K20004", f"=PRO.MACK.BOOT.SIMS({TRI},'Mack bootstrap'!$B$4,'Mack bootstrap'!$B$5,,,,{ORI})")
md.column_dimensions["N"].width = 30
for col in "ABCDEFGHIJKO":
    md.column_dimensions[col].width = 11

# A latest diagonal (a bordereau) projected with a selected pattern (decision 39).
dg = wb.create_sheet("Diagonal")
dg["A1"] = "prospicio XLL test: a latest diagonal projected with a selected pattern (PRO.DIAG.*)"; dg["A1"].font = Font(bold=True, size=13)
dg["A2"] = "Six accident years at 31 December 2024, factors to ultimate from a prior study, premium 200 each."; dg["A2"].font = note
for h, col in zip(["AY", "Age", "Latest", "Premium"], "ABCD"):
    dg[f"{col}3"] = h; dg[f"{col}3"].font = bold
for r, (y, age, x) in enumerate(zip(range(2019, 2025), [72, 60, 48, 36, 24, 12], [980, 860, 700, 520, 400, 150])):
    dg.cell(row=4 + r, column=1, value=y).font = bold
    for c, v in enumerate([age, x, 200]):
        dg.cell(row=4 + r, column=2 + c, value=v).fill = inp
dg["F3"] = "Pattern age"; dg["F4"] = "CDF"; dg["F3"].font = dg["F4"].font = bold
for k, (age, cdf) in enumerate(zip(range(12, 84, 12), [12 / 7, 1.3, 1.15, 1.07, 1.02, 1])):
    dg.cell(row=3, column=7 + k, value=age).font = bold
    dg.cell(row=4, column=7 + k, value=cdf).fill = inp
O_, A_, L_, P_, PAT, PAGES = "$A$4:$A$9", "$B$4:$B$9", "$C$4:$C$9", "$D$4:$D$9", "$G$4:$L$4", "$G$3:$L$3"
dg["A11"] = "Chain ladder"; dg["A11"].font = bold
dg["A12"] = ArrayFormula("A12:F19", f'=PRO.DIAG.CL({O_},{A_},{L_},{PAT},{PAGES},"cdf")')
dg["H11"] = "Bornhuetter-Ferguson, a priori 0.7"; dg["H11"].font = bold
dg["H12"] = ArrayFormula("H12:O19", f'=PRO.DIAG.BF({O_},{A_},{L_},{P_},0.7,{PAT},{PAGES},"cdf")')
dg["A21"] = "Benktander, 2 iterations"; dg["A21"].font = bold
dg["A22"] = ArrayFormula("A22:H29", f'=PRO.DIAG.BENKTANDER({O_},{A_},{L_},{P_},0.7,{PAT},2,{PAGES},"cdf")')
dg["J21"] = "Cape Cod, ages from the valuation date 31 December 2024"; dg["J21"].font = bold
dg["J22"] = ArrayFormula("J22:Q29", f'=PRO.DIAG.CAPECOD({O_},,{L_},{P_},{PAT},,,{PAGES},"cdf",,,45657)')
dg["A31"] = "Expected loss method, a priori 0.7"; dg["A31"].font = bold
dg["A32"] = ArrayFormula("A32:H39", f'=PRO.DIAG.EL({O_},{A_},{L_},{P_},0.7,{PAT},{PAGES},"cdf")')

# GenIns's diagonal (the Expected loss sheet's triangle) with its own volume-weighted
# factors to ultimate, against chainladder-python 0.10.1's DevelopmentConstant.
GENINS_CDF = [14.4465768672924, 4.138701016260354, 2.3685822129924685, 1.6251964813646196, 1.3844989689503577,
              1.2542756414472627, 1.1546635507833156, 1.0956368230729214, 1.017724725219544, 1.0]
dg["A42"] = "GenIns's latest diagonal with its own factors to ultimate, valued 31 December 2010"; dg["A42"].font = bold
dg["A43"] = "CDF"; dg["A43"].font = bold
for k, x in enumerate(GENINS_CDF):
    dg.cell(row=43, column=2 + k, value=x).fill = inp
dg["A45"] = ArrayFormula("A45:F56", "=PRO.DIAG.CL('Expected loss'!$A$4:$A$13,,PRO.TRI.LATEST('Expected loss'!$B$4:$K$13),$B$43:$K$43,,\"cdf\",,,40543)")

# Accident quarters of 2024 at December 2024, a pattern at 3, 12, 24 and 36 months.
dg["N42"] = "Accident quarters of 2024 (Excel dates), a pattern at 3, 12, 24, 36 months"; dg["N42"].font = bold
for h, col in zip(["Quarter", "Latest", "Pattern age", "CDF"], "NOQR"):
    dg[f"{col}43"] = h; dg[f"{col}43"].font = bold
for k, (m, x) in enumerate(zip([1, 4, 7, 10], [400, 300, 200, 100])):
    c = dg.cell(row=44 + k, column=14, value=datetime.date(2024, m, 1)); c.number_format = "yyyy-mm-dd"; c.fill = inp
    dg.cell(row=44 + k, column=15, value=x).fill = inp
for k, (age, cdf) in enumerate(zip([3, 12, 24, 36], [4, 2, 1.25, 1])):
    dg.cell(row=44 + k, column=17, value=age).fill = inp
    dg.cell(row=44 + k, column=18, value=cdf).fill = inp
dg["N50"] = ArrayFormula("N50:S55", '=PRO.DIAG.CL(N44:N47,,O44:O47,R44:R47,Q44:Q47,"cdf",,"quarter",45657)')

# Semiannual origins of 2023 and 2024 at December 2024 (24, 18 and 12
# months), the annual pattern read log-linearly at 18 months.
dg["W42"] = "Semiannual origins (Excel dates) at December 2024, CDF 2, 1.25, 1 at 12, 24, 36 months"; dg["W42"].font = bold
for h, col in zip(["Origin", "Latest", "Pattern age", "CDF"], ["W", "X", "Z", "AA"]):
    dg[f"{col}43"] = h; dg[f"{col}43"].font = bold
for k, (d, x) in enumerate(zip([datetime.date(2023, 1, 1), datetime.date(2023, 7, 1), datetime.date(2024, 1, 1)],
                               [500, 300, 100])):
    c = dg.cell(row=44 + k, column=23, value=d); c.number_format = "yyyy-mm-dd"; c.fill = inp
    dg.cell(row=44 + k, column=24, value=x).fill = inp
for k, (age, cdf) in enumerate(zip([12, 24, 36], [2, 1.25, 1])):
    dg.cell(row=44 + k, column=26, value=age).fill = inp
    dg.cell(row=44 + k, column=27, value=cdf).fill = inp
dg["W50"] = ArrayFormula("W50:AB54", '=PRO.DIAG.CL(W44:W46,,X44:X46,AA44:AA46,Z44:Z46,"cdf",,"semiannual",45657)')

# Past the pattern's end: nine factors of 1.1 from 12 to 108 months and a tail of 1.05.
dg["A59"] = "Past the pattern's end: factors 1.1 from 12 to 108 months, tail 1.05 spread one step per year"; dg["A59"].font = bold
for k, age in enumerate([144, 132, 120, 108, 12]):
    dg.cell(row=60 + k, column=1, value=2000 + k).font = bold
    dg.cell(row=60 + k, column=2, value=age).fill = inp
    dg.cell(row=60 + k, column=3, value=100).fill = inp
dg["E59"] = "LDF"; dg["E59"].font = bold
for k in range(9):
    dg.cell(row=59, column=6 + k, value=1.1).fill = inp
dg["A66"] = ArrayFormula("A66:F72", '=PRO.DIAG.CL(A60:A64,B60:B64,C60:C64,F59:N59,,,1.05,,,"tail")')

# An accident month's pattern t/12 averaged to accident years.
# Right of the past-end exhibit's LDF row (F59:N59), so no input lands in it.
dg["P59"] = "An accident month's t/12 at 1 to 12 months, averaged to accident years"; dg["P59"].font = bold
for k in range(12):
    dg.cell(row=60 + k, column=16, value=k + 1).fill = inp
    dg.cell(row=60 + k, column=17, value=(k + 1) / 12).fill = inp
dg["S60"] = ArrayFormula("S60:U83", '=PRO.PATTERN.AVERAGE(P60:P71,Q60:Q71,"month","year","percent")')

# The RAA sheet's chain ladder frozen in PRO.CL.LDF and applied to its own diagonal.
dg["A75"] = "RAA's own factors (the RAA sheet's PRO.CL.LDF row) applied to its latest diagonal"; dg["A75"].font = bold
dg["A76"] = ArrayFormula("A76:F87", "=PRO.DIAG.CL(RAA!$A$4:$A$13,,PRO.TRI.LATEST(RAA!$B$4:$K$13),RAA!$B$14:$J$14,,,RAA!$K$14,,DATE(1990,12,31))")

dg["H86"] = "Checks"; dg["H86"].font = bold
for h, col in zip(["Quantity", "prospicio", "Expected", "Match"], "HIJK"):
    dg[f"{col}87"] = h; dg[f"{col}87"].font = bold
for i, (name, f, x, tol) in enumerate([
        ("CL, 2024 ultimate (150 x 12/7)", "=E18", 257.142857142857, 1e-9),
        ("BF, 2024 ultimate", "=N18", 208.333333333333, 1e-9),
        ("Benktander, total ultimate", "=G29", 3930.086950714278, 1e-9),
        ("Cape Cod, total ultimate", "=P29", 4216.385925423543, 1e-9),
        ("EL, total ultimate", "=G39", 840, 1e-9),
        ("GenIns total, chainladder-python", "=E56", 53038945.61192431, 1e-9),
        ("Quarters, 2024Q1 ultimate", "=R51", 800, 1e-9),
        ("Quarters, 2024Q4 ultimate", "=R54", 400, 1e-9),
        ("Semiannual, 2023H2 ultimate (300 x sqrt(2.5))", "=AA52", 300 * 2.5 ** 0.5, 1e-9),
        ("132 months: CDF past the end", "=D68", 1.0253798721500054, 1e-12),
        ("Averaged: percent at 12 months (78/144)", "=T72", 78 / 144, 1e-12),
        ("RAA frozen: total ultimate against PRO.CL", "=E87", "=RAA!P14", 1e-12)]):
    row = 88 + i
    dg[f"H{row}"] = name; dg[f"I{row}"] = f; dg[f"J{row}"] = x
    dg[f"K{row}"] = f"=ABS(I{row}-J{row})<{tol}*MAX(1,ABS(J{row}))"
dg.column_dimensions["A"].width = 12
dg.column_dimensions["H"].width = 40
for col in [*"BCDEFGIJKLMOPQRSTUWXYZ", "AA", "AB"]:
    dg.column_dimensions[col].width = 12

# A Poisson GLM on validation/data/glm_policies.csv (600 policies) with an
# exposure offset, against statsmodels 0.15.0 (validation/reference/glm_statsmodels.csv).
REF = Path(__file__).resolve().parents[3] / "validation"
gl = wb.create_sheet("GLM")
gl["A1"] = "prospicio XLL test: a Poisson GLM with an exposure offset (PRO.GLM, .PREDICT, .RELATIVITIES)"; gl["A1"].font = Font(bold=True, size=13)
gl["A2"] = "600 policies from validation/data/glm_policies.csv; claims on age and region, log(exposure) as the offset."; gl["A2"].font = note
with open(REF / "data" / "glm_policies.csv") as f:
    policies = list(csv.DictReader(f))
for c, h in enumerate(["age", "region", "exposure", "claims"]):
    gl.cell(row=3, column=1 + c, value=h).font = bold
    for r, p in enumerate(policies):
        v = p[h]
        gl.cell(row=4 + r, column=1 + c, value=v if h == "region" else float(v)).fill = inp
DATA = f"$A$3:$D${3 + len(policies)}"
gl["F3"] = "GLM summary"; gl["F3"].font = bold
gl["F4"] = ArrayFormula("F4:L19", f'=PRO.GLM({DATA},"claims",{{"age","region"}},"poisson","log","exposure")')
gl["N3"] = "Relativities, 95%"; gl["N3"].font = bold
gl["N4"] = ArrayFormula("N4:R9", f'=PRO.GLM.RELATIVITIES({DATA},"claims",{{"age","region"}},"poisson","log","exposure")')
gl["F23"] = "New rows"; gl["F23"].font = bold
for c, h in enumerate(["age", "region", "exposure"]):
    gl.cell(row=24, column=6 + c, value=h).font = bold
for r, row in enumerate([(30, "A", 1), (45, "C", 0.5), (60, "D", 2)]):
    for c, v in enumerate(row):
        gl.cell(row=25 + r, column=6 + c, value=v).fill = inp
gl["I24"] = "Mean"; gl["I24"].font = bold
gl["I25"] = ArrayFormula("I25:I27", f'=PRO.GLM.PREDICT({DATA},"claims",{{"age","region"}},$F$24:$H$27,"poisson","log","exposure")')
gl["F30"] = "Check against statsmodels 0.15.0"; gl["F30"].font = bold
with open(REF / "reference" / "glm_statsmodels.csv") as f:
    sm = {(r["quantity"], r["term"]): float(r["expected"]) for r in csv.DictReader(f) if r["case"] == "poisson_log"}
for i, (name, cell) in enumerate([("(Intercept)", "H5"), ("age", "H6"), ("region[B]", "H7"),
                                  ("region[C]", "H8"), ("region[D]", "H9")]):
    row = 31 + i
    gl[f"F{row}"] = f"coefficient {name}"; gl[f"G{row}"] = f"={cell}"; gl[f"H{row}"] = sm[("coef", name)]
    gl[f"I{row}"] = f"=ABS(G{row}-H{row})<1e-8*MAX(1,ABS(H{row}))"
gl["F36"] = "deviance"; gl["G36"] = "=G11"; gl["H36"] = sm[("deviance", "")]
gl["I36"] = "=ABS(G36-H36)<1e-8*MAX(1,ABS(H36))"
gl.column_dimensions["F"].width = 24
for col in "GHIJKLNOPQR":
    gl.column_dimensions[col].width = 13

# Cash flows and returns beside Excel's own XNPV, XIRR and IRR on Microsoft's
# examples, and a collateralised deal through PRO.CASHFLOW.
cf = wb.create_sheet("Cash flows")
cf["A1"] = "prospicio XLL test: cash flows and returns (PRO.NPV, PRO.IRR, PRO.MOIC, PRO.CASHFLOW)"; cf["A1"].font = Font(bold=True, size=13)
cf["A2"] = "Flows on dates: PRO.NPV and PRO.IRR with unit \"dates\" are Excel's XNPV and XIRR."; cf["A2"].font = note
cf["A3"] = "Date"; cf["B3"] = "Flow"; cf["D3"] = "Yearly flow"; cf["A3"].font = cf["B3"].font = cf["D3"].font = bold
for r, (dt, x) in enumerate(zip([datetime.date(2008, 1, 1), datetime.date(2008, 3, 1), datetime.date(2008, 10, 30),
                                 datetime.date(2009, 2, 15), datetime.date(2009, 4, 1)],
                                [-10000, 2750, 4250, 3250, 2750])):
    cf.cell(row=4 + r, column=1, value=dt).fill = inp
    cf.cell(row=4 + r, column=2, value=x).fill = inp
for r, x in enumerate([-70000, 12000, 15000, 18000, 21000, 26000]):
    cf.cell(row=4 + r, column=4, value=x).fill = inp
checks = [
    ("PRO.NPV at 9% on dates", '=PRO.NPV(0.09,B4:B8,A4:A8,"dates")', "=XNPV(0.09,B4:B8,A4:A8)", "1e-9"),
    ("PRO.IRR on dates", '=PRO.IRR(B4:B8,A4:A8,"dates")', "=XIRR(B4:B8,A4:A8)", "1e-7"),
    ("PRO.IRR on five years", "=PRO.IRR(D4:D9)", "=IRR(D4:D9)", "1e-7"),
    ("PRO.IRR on four years", "=PRO.IRR(D4:D8)", "=IRR(D4:D8)", "1e-7"),
    ("PRO.NPV at 10% from time 0", "=PRO.NPV(0.1,D4:D9)", "=D4+NPV(0.1,D5:D9)", "1e-9"),
    ("PRO.MOIC(30, 40, 60)", "=PRO.MOIC(30,40,60)", 1.3, "1e-12"),
]
cf["F3"] = "Check"; cf["G3"] = "prospicio"; cf["H3"] = "Excel"; cf["I3"] = "Agrees"
for c in "FGHI":
    cf[f"{c}3"].font = bold
for i, (name, ours, excel, tol) in enumerate(checks):
    row = 4 + i
    cf[f"F{row}"] = name; cf[f"G{row}"] = ours; cf[f"H{row}"] = excel
    cf[f"I{row}"] = f"=ABS(G{row}-H{row})<{tol}*MAX(1,ABS(H{row}))"

cf["A12"] = "A collateralised deal: premium 1,000 written at once, earned over 12 months, paid over 30, commuting after 36"; cf["A12"].font = bold
cf["A13"] = "Terms"; cf["D13"] = "Requirement"; cf["G13"] = "At commutation"
for c in "ADG":
    cf[f"{c}13"].font = bold
terms = [("premium", 1000), ("commutation_month", 36), ("collected_rate", 0.9), ("annual_yield", 0.04),
         ("opening_fixed", 50), ("expense_rate", 0.01), ("collateral_release", "quarterly"),
         ("equity_release", "at-commutation"), ("fee_breakpoints", 0, 100), ("fee_slopes", 0.1, 0.2)]
for r, (label, *values) in enumerate(terms):
    cf.cell(row=14 + r, column=1, value=label)
    for c, v in enumerate(values):
        cf.cell(row=14 + r, column=2 + c, value=v).fill = inp
for r, (label, *values) in enumerate([("outstanding", 1), ("unearned-premium", 0.3), ("floor", 20)]):
    cf.cell(row=14 + r, column=4, value=label)
    cf.cell(row=14 + r, column=5, value=values[0]).fill = inp
cf["G14"] = "outstanding"; cf["H14"] = 1.1; cf["H14"].fill = inp
cf["J13"] = "Month"; cf["K13"] = "written"; cf["L13"] = "earned"; cf["M13"] = "paid"
for c in "JKLM":
    cf[f"{c}13"].font = bold
for t in range(40):
    row = 14 + t
    cf.cell(row=row, column=10, value=t)
    for c, v in enumerate([1, min((t + 1) / 12, 1), min((t + 1) / 30, 1)]):
        cf.cell(row=row, column=11 + c, value=v).fill = inp
cf["O13"] = "Loss ratio"; cf["O13"].font = bold
for r, u in enumerate([0.4, 0.65, 0.9, 1.3, 2.0]):
    cf.cell(row=14 + r, column=15, value=u).fill = inp
DEAL = "O14:O18,A14:C23,K13:M53,D14:E16"
cf["A26"] = "Per simulation, with the NPV at 8%"; cf["A26"].font = bold
cf["A27"] = ArrayFormula("A27:K32", f'=PRO.CASHFLOW({DEAL},,"simulations",0.08,G14:H14)')
cf["A35"] = "The first simulation's accounts"; cf["A35"].font = bold
cf["A36"] = ArrayFormula("A36:O76", f'=PRO.CASHFLOW({DEAL},,"trace",,G14:H14)')
cf["Q26"] = "Check"; cf["Q26"].font = bold
cf["Q27"] = "money is conserved"; cf["R27"] = "=ABS(B76)<1e-6"
cf["Q28"] = "NPV of the trace's flows at 8%"; cf["R28"] = '=ABS(PRO.NPV(0.08,N37:N72,J14:J49,"months")-C28)<1e-6'
cf.column_dimensions["A"].width = 20
cf.column_dimensions["F"].width = 28
cf.column_dimensions["Q"].width = 34

# The P&L ledger: aggregate 1.0.1's Example 3 (annual losses 2 to 8, premium
# 10, a 10% premium expense, an aggregate cover of 3 xs 5 for 1 with a 10%
# commission), against its printed sheet (validation/reference/pnl_builders_aggregate.csv).
pl = wb.create_sheet("P&L")
pl["A1"] = "prospicio XLL test: the P&L ledger (PRO.PNL), aggregate's Example 3"; pl["A1"].font = Font(bold=True, size=13)
pl["A2"] = "Annual gross losses and their probabilities; the book and the cover it buys."; pl["A2"].font = note
pl["A3"] = "Loss"; pl["B3"] = "Prob"; pl["A3"].font = pl["B3"].font = bold
for r, (x, p) in enumerate(zip([2, 3, 4, 5, 6, 8], [0.25, 0.3, 0.09, 0.2, 0.12, 0.04])):
    pl.cell(row=4 + r, column=1, value=x).fill = inp
    pl.cell(row=4 + r, column=2, value=p).fill = inp
pl["D3"] = "Premium"; pl["E3"] = 10; pl["D3"].font = bold; pl["E3"].fill = inp
for c, h in enumerate(["label", "premium", "loss", "fixed"]):
    pl.cell(row=5, column=4 + c, value=h).font = bold
pl["D6"] = "Acquisition"; pl["E6"] = 0.1; pl["E6"].fill = inp
for c, (h, x) in enumerate(zip(["name", "limit", "attachment", "premium", "commission"], ["agg 3 xs 5", 3, 5, 1, 0.1])):
    pl.cell(row=8, column=4 + c, value=h).font = bold
    pl.cell(row=9, column=4 + c, value=x).fill = inp
ARGS = "$A$4:$A$9,$B$4:$B$9,$E$3,$D$5:$G$6,$D$8:$H$9"
pl["A12"] = "The walk: the gross book, then the cover"; pl["A12"].font = bold
pl["A13"] = ArrayFormula("A13:P28", f'=PRO.PNL({ARGS},"walk")')
pl["A31"] = "The card"; pl["A31"].font = bold
pl["A32"] = ArrayFormula("A32:I43", f'=PRO.PNL({ARGS},"walk","card")')
pl["K31"] = "The ratios"; pl["K31"].font = bold
pl["K32"] = ArrayFormula("K32:X35", f'=PRO.PNL({ARGS},"walk","ratios")')
pl["A46"] = "Check against aggregate 1.0.1"; pl["A46"].font = bold
with open(REF / "reference" / "pnl_builders_aggregate.csv") as f:
    ex3 = {(r["row"], r["col"]): float(r["expected"]) for r in csv.DictReader(f)
           if r["case"] == "ex3a" and r["quantity"] == "sheet"}
for i, (step, side, line, col, idx) in enumerate([
        ("All", "Margin", "Net", "EX", 3), ("agg 3 xs 5", "Obligation", "agg 3 xs 5 recovery", "EX", 3),
        ("All", "Margin", "Net", "k01", 7), ("All", "Margin", "Net", "k99", 15)]):
    row = 47 + i
    key = f"{step}|{side}|{line}"
    col_letter = "ABCDEFGHIJKLMNOP"[idx]
    pl[f"A{row}"] = f"{key} {col}"
    pl[f"B{row}"] = (f'=SUMPRODUCT(($A$14:$A$28="{step}")*($B$14:$B$28="{side}")*($C$14:$C$28="{line}"),'
                     f'${col_letter}$14:${col_letter}$28)')
    pl[f"C{row}"] = ex3[(key, col)]
    pl[f"D{row}"] = f"=ABS(B{row}-C{row})<1e-9"
pl.column_dimensions["A"].width = 30
pl.column_dimensions["C"].width = 22

# Catastrophe model output: Alder Mutual's year and event loss tables
# (examples/data/cat_book_*.csv), against the numbers the guide
# docs/guides/cat-2-risk-metrics.md prints.
ct = wb.create_sheet("Cat")
ct["A1"] = "prospicio XLL test: year and event loss tables (PRO.YLT.STATS, PRO.YLT.OEP, PRO.ELT.STATS, PRO.ELT.OEP)"
ct["A1"].font = Font(bold=True, size=13)
ct["A2"] = "Alder Mutual's 2,000 simulated years and its event loss table, as in the guide cat-2-risk-metrics."; ct["A2"].font = note
for c, h in enumerate(["Year", "Loss", "Time"]):
    ct.cell(row=3, column=1 + c, value=h).font = bold
with open(REF.parent / "examples" / "data" / "cat_book_yelt.csv") as f:
    alder = [r for r in csv.DictReader(f) if r["cedent"] == "alder"]
for r, x in enumerate(alder):
    ct.cell(row=4 + r, column=1, value=int(x["year"])).fill = inp
    ct.cell(row=4 + r, column=2, value=float(x["loss"])).fill = inp
    ct.cell(row=4 + r, column=3, value=f"=(({int(x['day'])}-1)*24+{int(x['hour'])})/8760")
last = 3 + len(alder)
for c, h in enumerate(["Rate", "Mean", "SD independent", "SD correlated", "Exposure"]):
    ct.cell(row=3, column=5 + c, value=h).font = bold
with open(REF.parent / "examples" / "data" / "cat_book_elt_alder.csv") as f:
    events = list(csv.DictReader(f))
for r, x in enumerate(events):
    for c, k in enumerate(["rate", "mean", "sd_independent", "sd_correlated", "exposure"]):
        ct.cell(row=4 + r, column=5 + c, value=float(x[k])).fill = inp
elast = 3 + len(events)
Y = f"$A$4:$A${last},$B$4:$B${last}"
E = f"$E$4:$E${elast},$F$4:$F${elast},$G$4:$G${elast},$H$4:$H${elast},$I$4:$I${elast}"
T = f"$C$4:$C${last}"
ct["K3"] = "Annual totals"; ct["K3"].font = bold
ct["K4"] = ArrayFormula("K4:L19", f"=PRO.YLT.STATS({Y},2000,{{0.9,0.99}})")
ct["N3"] = "Amounts"; ct["N3"].font = bold
ct["N4"] = 30; ct["N5"] = 70; ct["N6"] = 110
for c in ("N4", "N5", "N6"):
    ct[c].fill = inp
ct["P3"] = "Exceedance at amounts, raw"; ct["P3"].font = bold
ct["P4"] = ArrayFormula("P4:T7", f"=PRO.YLT.OEP({Y},,N4:N6)")
ct["P10"] = "PMLs by return period, a 72-hour clause"; ct["P10"].font = bold
ct["N11"] = 10; ct["N12"] = 50; ct["N13"] = 100; ct["N14"] = 250
for c in ("N11", "N12", "N13", "N14"):
    ct[c].fill = inp
ct["P11"] = ArrayFormula("P11:U15", f"=PRO.YLT.OEP({Y},N11:N14,,,,{T},72)")
ct["K22"] = "The event loss table"; ct["K22"].font = bold
ct["K23"] = ArrayFormula("K23:L29", f"=PRO.ELT.STATS({E})")
ct["P18"] = "The event loss table's OEP"; ct["P18"].font = bold
ct["P19"] = ArrayFormula("P19:T23", f"=PRO.ELT.OEP({E},N11:N14)")
ct["P26"] = "Check against the guide"; ct["P26"].font = bold
checks = [
    ("Average annual loss", "L9", 16.778, 3), ("SD", "L10", 35.158, 3), ("VaR 99%", "L17", 176.824, 3),
    ("TVaR 99%", "L19", 255.969, 3), ("OEP at 30", "Q5", 0.1070, 4), ("AEP at 110", "R7", 0.0235, 4),
    ("100-year OEP loss, 72 hours", "R14", 163.418, 3), ("250-year AEP loss", "S15", 232.741, 3),
    ("ELT events a year", "L24", 2.3365, 4), ("ELT expected annual loss", "L26", 16.056, 3),
    ("ELT 100-year occurrence loss", "R22", 147.938, 3),
]
for i, (name, cell, value, digits) in enumerate(checks):
    row = 27 + i
    ct[f"P{row}"] = name
    ct[f"Q{row}"] = f"={cell}"
    ct[f"R{row}"] = value
    ct[f"S{row}"] = f"=ROUND(Q{row},{digits})=R{row}"
ct.column_dimensions["K"].width = 20
ct.column_dimensions["P"].width = 30

# Capital allocation: Pricing Insurance Risk's Simple Discrete Example
# (Exercises 205 and 206) and numpy's allocations of four units
# (validation/reference/allocation_numpy.csv).
al = wb.create_sheet("Allocation")
al["A1"] = "prospicio XLL test: capital allocation (PRO.ALLOCATE, PRO.ALLOCATE.TOTAL)"; al["A1"].font = Font(bold=True, size=13)
al["A2"] = "Pricing Insurance Risk's Simple Discrete Example: two units over 16 equally likely years."; al["A2"].font = note
al["A3"] = "X1"; al["B3"] = "X2"; al["A3"].font = al["B3"].font = bold
r = 4
for a, k in [(0, 2), (8, 1), (10, 1)]:
    for b, j in [(0, 2), (1, 1), (90, 1)]:
        for _ in range(k * j):
            al.cell(row=r, column=1, value=a).fill = inp
            al.cell(row=r, column=2, value=b).fill = inp
            r += 1
al["D3"] = "A premium of 80 by stand-alone TVaR at 75% (Exercise 205)"; al["D3"].font = bold
al["D4"] = ArrayFormula("D4:G8", '=PRO.ALLOCATE.TOTAL(A3:B19,80,"haircut","tvar",0.75)')
al["D11"] = "At the equal risk level (Exercise 206)"; al["D11"].font = bold
al["D12"] = ArrayFormula("D12:G16", '=PRO.ALLOCATE.TOTAL(A3:B19,80,"equal_risk")')
al["D19"] = "TVaR 90% by Euler"; al["D19"].font = bold
al["D20"] = ArrayFormula("D20:H24", '=PRO.ALLOCATE(A3:B19,"tvar",0.9)')
al["J3"] = "Four units over 400 simulations (numpy reference)"; al["J3"].font = bold
for c, h in enumerate(["U1", "U2", "U3", "U4"]):
    al.cell(row=4, column=10 + c, value=h).font = bold
for i in range(400):
    a = (i * 37) % 101
    for c, v in enumerate([a, (i * 53) % 97 + a // 2, (i * i) % 89, 3 * ((i * 29) % 41) + (a // 10) ** 2]):
        al.cell(row=5 + i, column=10 + c, value=v).fill = inp
al["O3"] = "Shapley under PH 0.5"; al["O3"].font = bold
al["O4"] = ArrayFormula("O4:S10", '=PRO.ALLOCATE(J4:M404,"ph",0.5,"shapley")')
al["O13"] = "Covariance under dual power 3"; al["O13"].font = bold
al["O14"] = ArrayFormula("O14:S20", '=PRO.ALLOCATE(J4:M404,"dual",3,"covariance")')
with open(REF / "reference" / "allocation_numpy.csv") as f:
    ref = {(r["distortion"], r["arg"], r["method"], r["component"]): float(r["expected"]) for r in csv.DictReader(f)}
al["O23"] = "Check"; al["P23"] = "prospicio"; al["Q23"] = "Expected"; al["R23"] = "Agrees"
for c in "OPQR":
    al[f"{c}23"].font = bold
checks = [
    ("Exercise 205: X1 gets 8", "F5", 8.0, "1e-12"), ("Exercise 205: X2 gets 72", "F6", 72.0, "1e-12"),
    ("Exercise 206: the level 0.6796", "E16", 0.6796, "5e-5"), ("Exercise 206: X1 gets 9.56", "F13", 9.56, "5e-3"),
    ("Shapley, PH 0.5, unit 1", "Q5", ref[("proportional_hazard", "0.5", "shapley", "0")], "1e-9"),
    ("Shapley, PH 0.5, unit 4", "Q8", ref[("proportional_hazard", "0.5", "shapley", "3")], "1e-9"),
    ("Dual power 3, the portfolio", "Q20", ref[("dual_power", "3.0", "total", "")], "1e-9"),
    ("Covariance, dual power 3, unit 2", "Q16", ref[("dual_power", "3.0", "covariance", "1")], "1e-9"),
]
for i, (name, cell, value, tol) in enumerate(checks):
    row = 24 + i
    al[f"O{row}"] = name; al[f"P{row}"] = f"={cell}"; al[f"Q{row}"] = value
    al[f"R{row}"] = f"=ABS(P{row}-Q{row})<{tol}*MAX(1,ABS(Q{row}))"
al.column_dimensions["D"].width = 14
al.column_dimensions["O"].width = 34

wb.save(out)

METADATA = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" \
xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">\
<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" \
pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" \
clearComments="1" assign="1" coerce="1" cellMeta="1"/></metadataTypes>\
<futureMetadata name="XLDAPR" count="1"><bk><extLst>\
<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">\
<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>\
<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>"""
REL = ('<Relationship Id="rIdMeta" Target="metadata.xml" Type="http://schemas.openxmlformats.org/'
       'officeDocument/2006/relationships/sheetMetadata"/>')
OVERRIDE = ('<Override PartName="/xl/metadata.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/>')


def make_dynamic(path):
    """Turns every legacy array formula in `path` into a spilling one."""
    with zipfile.ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    n = 0
    for name in [p for p in parts if p.startswith("xl/worksheets/sheet")]:
        sheet, k = re.subn(
            r'<c (r="[A-Z]+[0-9]+")([^>]*)><f t="array"', r'<c \1\2 cm="1"><f t="array"', parts[name].decode()
        )
        parts[name] = sheet.encode()
        n += k
    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    parts["xl/_rels/workbook.xml.rels"] = rels.replace("</Relationships>", REL + "</Relationships>").encode()
    types = parts["[Content_Types].xml"].decode()
    parts["[Content_Types].xml"] = types.replace("</Types>", OVERRIDE + "</Types>").encode()
    parts["xl/metadata.xml"] = METADATA.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return n


print("wrote", out, "with", make_dynamic(out), "spilling formulas")
